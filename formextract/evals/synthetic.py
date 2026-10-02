"""Synthetic golden-set fixtures: the spec's layout fragment as PDF and XLSX."""
from __future__ import annotations

from pathlib import Path
from typing import Callable

import pymupdf

ROWS_PDF = [
    (62.0, [(50.0, "VENDOR COMPLIANCE CHECKLIST", 14.0)]),
    (
        100.0,
        [
            (50.0, "Which US states do they ship to", 10.0),
            (380.0, "X", 10.0),
            (394.0, "All", 10.0),
        ],
    ),
    (
        116.0,
        [
            (70.0, "if not all please check those that apply", 9.0),
        ]
    ),
    (116.0, [(380.0 + i * 18.0, s, 9.0) for i, s in enumerate(
        "AK AL AR AZ CA CO CT DE FL GA".split())]),
    (
        132.0,
        [(380.0 + i * 18.0, s, 9.0) for i, s in enumerate(
            "HI IA ID IL IN KS KY LA MA MD".split())],
    ),
    (
        160.0,
        [
            (50.0, "Do they ship to Canada?", 10.0),
            (340.0, "Yes", 10.0),
            (376.0, "X", 10.0),
            (390.0, "No", 10.0),
        ],
    ),
    (
        184.0,
        [
            (50.0, "Vendor Status", 10.0),
            (300.0, "Approved", 10.0),
            (356.0, "X", 10.0),
            (368.0, "Provisional", 10.0),
            (462.0, "Confirm on vendor portal", 9.0),
        ],
    ),
    (
        208.0,
        [
            (50.0, "Order Log", 10.0),
            (300.0, "X", 10.0),
            (314.0, "Support Order Log (Check if the vendor is exporting its order log directly)", 8.0),
        ],
    ),
]

_STATE_LINE_1 = "AK AL AR AZ CA CO CT DE FL GA"
_STATE_LINE_2 = "HI IA ID IL IN KS KY LA MA MD"


def build_spec_fragment_pdf(dest: Path) -> Path:
    dest.mkdir(parents=True, exist_ok=True)
    path = dest / "spec_fragment.pdf"
    if path.exists():
        path.unlink()
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    for _y, spans in ROWS_PDF:
        for x, text, size in spans:
            page.insert_text((x, _y), text, fontname="helv", fontsize=size)
    doc.save(str(path))
    doc.close()
    return path


def build_spec_fragment_xlsx(dest: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font

    dest.mkdir(parents=True, exist_ok=True)
    path = dest / "spec_fragment.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Checklist"
    ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
    ws["A1"].font = Font(bold=True, sz=14)

    ws["A3"] = "Which US states do they ship to"
    ws["F3"] = "X"
    ws["G3"] = "All"

    ws["A4"] = "if not all please check those that apply"
    ws["F4"] = _STATE_LINE_1
    ws["F5"] = _STATE_LINE_2

    ws["A7"] = "Do they ship to Canada?"
    ws["F7"] = "Yes"
    ws["G7"] = "X"
    ws["H7"] = "No"

    ws["A9"] = "Vendor Status"
    ws["F9"] = "Approved"
    ws["G9"] = "X"
    ws["H9"] = "Provisional"
    ws["J9"] = "Confirm on vendor portal"

    ws["A11"] = "Order Log"
    ws.merge_cells("F11:K11")
    ws["F11"] = "X Support Order Log (Check if the vendor is exporting its order log directly)"

    wb.save(str(path))
    wb.close()
    return path


FACTORIES: dict[str, Callable[[Path], Path]] = {
    "spec_fragment_pdf": build_spec_fragment_pdf,
    "spec_fragment_xlsx": build_spec_fragment_xlsx,
}


def build(name: str, dest_dir: Path) -> Path:
    if name not in FACTORIES:
        raise KeyError(f"unknown synthetic fixture: {name}")
    return FACTORIES[name](dest_dir)
