"""Native .xlsx ingestion → Element stream (grid coordinates, merged cells, fills)."""
from __future__ import annotations

from pathlib import Path
from typing import Callable

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string

from ..model import BBox, Element


def _argb_to_int(color) -> int | None:
    rgb = getattr(color, "rgb", None)
    if isinstance(rgb, str) and len(rgb) in (6, 8):
        try:
            return int(rgb[-6:], 16)
        except ValueError:
            return None
    return None


def _column_geometry(ws) -> tuple[Callable[[int], float], Callable[[int], float]]:
    """Width-aware x coordinates shared by the merged and unmerged bbox
    branches: `col_x0(c)` / `col_x1(c)` give the cumulative width up to
    (excluding / including) column `c`, so x is never a mix of raw column
    index and width units.

    Declared widths come from `ws.column_dimensions`, iterating `.values()`
    and expanding `range(dim.min, dim.max + 1)` because a range-defined
    width (`<col min="4" max="8" width="20"/>`) is one entry keyed by the
    first letter only — a per-letter lookup would silently keep index math
    for the whole block. A `None` width (style-only `<col>` entry) uses the
    sheet default; with no declaration at all each column counts as 1.0,
    which makes cumulative width equal raw column index on width-less
    sheets (the absolute bbox pins in tests/test_ingest.py depend on that).
    Excel's rendered default (~8.43) is deliberately NOT assumed.

    Known edge case: hidden columns report `width=0`, which collapses two
    genuinely separated columns into adjacent x — not absorbed silently.

    y stays raw row-index units. Safe only because layout.py never compares
    x and y in the same expression (`_segment_columns` is x-only; banding's
    `tol_y` is height-based/y-only) — a future area/aspect-ratio heuristic
    must convert units first rather than mixing them.
    """
    fmt = getattr(ws, "sheet_format", None)
    default = float(fmt.defaultColWidth) if getattr(fmt, "defaultColWidth", None) else 1.0
    widths: dict[int, float] = {}
    for dim in ws.column_dimensions.values():
        w = float(dim.width) if dim.width is not None else default
        lo, hi = dim.min, dim.max
        if lo is None or hi is None:
            idx = column_index_from_string(str(dim.index))
            lo = hi = idx
        for c in range(int(lo), int(hi) + 1):
            widths[c] = w
    cum: dict[int, float] = {0: 0.0}
    highest = 0

    def _upto(col: int) -> float:
        nonlocal highest
        while highest < col:
            highest += 1
            cum[highest] = cum[highest - 1] + widths.get(highest, default)
        return cum[col]

    def col_x0(col: int) -> float:
        return _upto(col - 1)

    def col_x1(col: int) -> float:
        return _upto(col)

    return col_x0, col_x1


def ingest_xlsx(path: str | Path) -> tuple[list[Element], list[str], dict[str, str]]:
    wb = openpyxl.load_workbook(filename=str(path), data_only=True)
    elements: list[Element] = []
    sheet_names = list(wb.sheetnames)
    sheet_state = {
        ws.title: getattr(ws, "sheet_state", "visible") for ws in wb.worksheets
    }
    for sheet_index, ws in enumerate(wb.worksheets):
        col_x0, col_x1 = _column_geometry(ws)
        merged_anchors: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        merged_members: set[tuple[int, int]] = set()
        for rng in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = (
                rng.min_col,
                rng.min_row,
                rng.max_col,
                rng.max_row,
            )
            merged_anchors[(min_row, min_col)] = (min_col, min_row, max_col, max_row)
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    if (r, c) != (min_row, min_col):
                        merged_members.add((r, c))
        for row in ws.iter_rows():
            for cell in row:
                rc = (cell.row, cell.column)
                if isinstance(cell, MergedCell) or rc in merged_members:
                    continue
                if cell.value is None:
                    continue
                if rc in merged_anchors:
                    min_col, min_row, max_col, max_row = merged_anchors[rc]
                    x0, y0 = col_x0(min_col), min_row - 1
                    x1, y1 = col_x1(max_col), max_row
                    merged = True
                else:
                    x0, y0 = col_x0(cell.column), cell.row - 1
                    x1, y1 = col_x1(cell.column), cell.row
                    merged = False
                font = cell.font
                elements.append(
                    Element(
                        element_id=f"{ws.title}!{cell.row}:{cell.column}",
                        text=str(cell.value),
                        bbox=BBox(page=sheet_index, x0=x0, y0=y0, x1=x1, y1=y1),
                        font=getattr(font, "name", None),
                        size=float(font.sz) if getattr(font, "sz", None) else None,
                        color=_argb_to_int(getattr(font, "color", None)),
                        fill=_argb_to_int(getattr(cell.fill, "start_color", None)),
                        sheet=ws.title,
                        merged=merged,
                    )
                )
    wb.close()
    return elements, sheet_names, sheet_state
