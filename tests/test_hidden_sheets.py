from __future__ import annotations

import json
import re
from pathlib import Path

from formextract.ingest import ingest
from formextract.model import InstanceStatus
from formextract.pipeline import Pipeline, PipelineConfig
from formextract.resolve import LLMResponse
from formextract.store import Store


def _hidden_workbook(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    visible = wb.active
    visible.title = "Visible"
    visible["A1"] = "Order Log"
    visible["F1"] = "X Support Order Log"

    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "Order Log"
    hidden["F1"] = "X Support Order Log"
    hidden.sheet_state = "hidden"

    very = wb.create_sheet("VeryHidden")
    very["A1"] = "Order Log"
    very.sheet_state = "veryHidden"

    wb.save(path)
    return path


def test_xlsx_captures_sheet_state(tmp_path: Path):
    path = _hidden_workbook(tmp_path / "hidden.xlsx")
    result = ingest(path)
    assert result.sheet_names == ["Visible", "Hidden", "VeryHidden"]
    assert result.sheet_state == {
        "Visible": "visible",
        "Hidden": "hidden",
        "VeryHidden": "veryHidden",
    }


class _TabClient:
    def complete(self, prompt, *, model, params):
        match = re.search(r"Layout projection for `([^`]+)`", prompt)
        tab = match.group(1) if match else "?"
        return LLMResponse(
            text=json.dumps(
                {
                    "fields": [
                        {
                            "label": f"Field on {tab}",
                            "control_type": "text",
                            "answer": [tab],
                            "annotations": [],
                        }
                    ]
                }
            ),
            model=model,
            params=params,
        )


def test_hidden_sheets_flagged_and_excluded(tmp_path: Path):
    path = _hidden_workbook(tmp_path / "hidden.xlsx")
    store = Store(tmp_path / "store")
    record = Pipeline(store, _TabClient(), PipelineConfig()).run(path)

    assert record.status is InstanceStatus.COMPLETE
    assert record.source.sheet_state == {
        "Visible": "visible",
        "Hidden": "hidden",
        "VeryHidden": "veryHidden",
    }
    assert record.hidden_sheets == ["Hidden", "VeryHidden"]
    assert record.tabs == ["Visible"]
    assert [f.label_text for f in record.fields] == ["Field on Visible"]
    # hidden-sheet elements never reach layout/fields
    for region in record.regions:
        assert not any(eid.startswith(("Hidden!", "VeryHidden!")) for eid in region.element_ids)


def test_include_hidden_sheets_flag(tmp_path: Path):
    path = _hidden_workbook(tmp_path / "hidden.xlsx")
    store = Store(tmp_path / "store")
    record = Pipeline(
        store, _TabClient(), PipelineConfig(include_hidden_sheets=True)
    ).run(path)

    assert record.hidden_sheets == []
    assert record.tabs == ["Visible", "Hidden", "VeryHidden"]
    assert [f.label_text for f in record.fields] == [
        "Field on Visible",
        "Field on Hidden",
        "Field on VeryHidden",
    ]


def test_include_hidden_sheets_changes_cache_key(tmp_path: Path):
    path = _hidden_workbook(tmp_path / "hidden.xlsx")
    store = Store(tmp_path / "store")
    excluded = Pipeline(store, _TabClient(), PipelineConfig()).run(path)
    included = Pipeline(
        store, _TabClient(), PipelineConfig(include_hidden_sheets=True)
    ).run(path)
    assert excluded.instance_id != included.instance_id
