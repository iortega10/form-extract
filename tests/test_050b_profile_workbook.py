from __future__ import annotations

import importlib
import json
import re
import sys
from pathlib import Path

from formextract.pipeline import Pipeline, PipelineConfig
from formextract.resolve import LLMResponse
from formextract.store import Store

TOOL_PATH = Path(__file__).resolve().parents[1] / "tools" / "profile_workbook.py"


def _load_tool(monkeypatch):
    monkeypatch.syspath_prepend(str(TOOL_PATH.parent))
    return importlib.import_module("profile_workbook")


def _workbook(path: Path) -> Path:
    """First sheet hidden, two visible sheets with the same geometry."""
    from openpyxl import Workbook

    wb = Workbook()
    hidden = wb.active
    hidden.title = "HiddenFirst"
    hidden["A1"] = "Order Log"
    hidden.sheet_state = "hidden"

    one = wb.create_sheet("VisibleOne")
    one["A1"] = "Alpha Log"

    two = wb.create_sheet("VisibleTwo")
    two["A1"] = "Order Log"

    wb.save(path)
    wb.close()
    return path


class _RecordingClient:
    def __init__(self) -> None:
        self.tabs: list[str] = []

    def complete(self, prompt: str, *, model: str, params: dict) -> LLMResponse:
        self.tabs.append(
            re.search(r"Layout projection for `([^`]+)`", prompt).group(1)
        )
        return LLMResponse(
            text=json.dumps(
                {
                    "fields": [
                        {
                            "label": "L",
                            "control_type": "text",
                            "options": [],
                            "answer": ["v"],
                            "annotations": [],
                        }
                    ]
                }
            ),
            model=model,
            params=params,
            tokens=1,
            latency_ms=0,
        )


def _pipeline_tabs(path: Path, store_dir: Path, *, include_hidden_sheets=False):
    client = _RecordingClient()
    Pipeline(
        Store(store_dir),
        client,
        PipelineConfig(include_hidden_sheets=include_hidden_sheets),
    ).run(path)
    return client.tabs


def test_profile_tool_chunk_keys_equal_pipeline_chunk_keys(tmp_path: Path, monkeypatch):
    path = _workbook(tmp_path / "wb.xlsx")
    tool = _load_tool(monkeypatch)

    profiled = [chunk.key for chunk in tool.profile(path).chunks]
    authored = _pipeline_tabs(path, tmp_path / "store-default")

    assert profiled == authored
    assert authored == ["VisibleOne", "VisibleTwo"]


def test_profile_tool_include_hidden_matches_pipeline(tmp_path: Path, monkeypatch):
    path = _workbook(tmp_path / "wb.xlsx")
    tool = _load_tool(monkeypatch)

    profiled = [chunk.key for chunk in tool.profile(path, include_hidden=True).chunks]
    authored = _pipeline_tabs(
        path, tmp_path / "store-hidden", include_hidden_sheets=True
    )

    assert profiled == authored
    assert authored == ["HiddenFirst", "VisibleOne", "VisibleTwo"]


def test_profile_tool_prints_chunk_count_and_reuse_misses_without_names(
    tmp_path: Path, monkeypatch, capsys
):
    path = _workbook(tmp_path / "wb.xlsx")
    tool = _load_tool(monkeypatch)

    monkeypatch.setattr(sys, "argv", ["profile_workbook.py", str(path), "--no-names"])
    tool.main()
    out = capsys.readouterr().out

    assert "chunks authored: 2" in out
    for name in ("HiddenFirst", "VisibleOne", "VisibleTwo", "Alpha Log", "Order Log"):
        assert name not in out
    assert "reuse misses (equal geometry, differing anchors): 1" in out
    assert "page 1 vs page 2: anchors_symmetric_difference=2 jaccard=0.000" in out

    monkeypatch.setattr(sys, "argv", ["profile_workbook.py", str(path)])
    tool.main()
    named = capsys.readouterr().out

    assert "chunks authored: 2" in named
    assert "VisibleOne: prompt_chars" in named
    assert "reuse misses (equal geometry, differing anchors): 1" in named
    assert "VisibleOne vs VisibleTwo: anchors_symmetric_difference=2" in named
