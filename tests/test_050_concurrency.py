from __future__ import annotations

import json
import re
import threading
import time
from pathlib import Path

import pytest

from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import ElementRef, InstanceStatus, ProvenanceSource
from formextract.pipeline import Pipeline, PipelineConfig
from formextract.resolve import LLMResponse, ProjectionChunk, author_drafts
from formextract.store import Store


def _valid_response(label):
    return json.dumps(
        {
            "fields": [
                {
                    "label": label,
                    "control_type": "text",
                    "answer": ["v"],
                    "annotations": [],
                }
            ]
        }
    )


def _chunks(n):
    return [ProjectionChunk(key=f"tab{i}", text=f"chunk {i}") for i in range(n)]


def test_serial_equals_workers_for_author_drafts(store):
    chunks = _chunks(5)

    class Client:
        def complete(self, prompt, *, model, params):
            match = re.search(r"Layout projection for `([^`]+)`", prompt)
            tab = match.group(1) if match else "?"
            time.sleep(0.01)
            return LLMResponse(text=_valid_response(tab), model=model, params=params)

    serial_drafts, serial_calls, serial_errors = author_drafts(
        chunks,
        Client(),
        model="m",
        params={"temperature": 0},
        store=store,
        max_workers=1,
    )
    parallel_drafts, parallel_calls, parallel_errors = author_drafts(
        chunks,
        Client(),
        model="m",
        params={"temperature": 0},
        store=store,
        max_workers=4,
    )

    assert [d.label for d in serial_drafts] == [d.label for d in parallel_drafts]
    assert [d.control_type for d in serial_drafts] == [
        d.control_type for d in parallel_drafts
    ]
    assert [c.prompt_hash for c in serial_calls] == [
        c.prompt_hash for c in parallel_calls
    ]
    assert [c.response_ref for c in serial_calls] == [
        c.response_ref for c in parallel_calls
    ]
    assert serial_errors == parallel_errors


def test_workers_overlap_proven_by_barrier(store):
    n = 4
    barrier = threading.Barrier(n)
    chunks = _chunks(n)

    class Client:
        def __init__(self):
            self.in_flight = 0
            self.max_in_flight = 0
            self.lock = threading.Lock()

        def complete(self, prompt, *, model, params):
            with self.lock:
                self.in_flight += 1
                self.max_in_flight = max(self.max_in_flight, self.in_flight)
            barrier.wait(timeout=10)
            with self.lock:
                self.in_flight -= 1
            return LLMResponse(text=_valid_response("F"), model=model, params=params)

    client = Client()
    author_drafts(
        chunks,
        client,
        model="m",
        params={"temperature": 0},
        store=store,
        max_workers=4,
    )
    assert client.max_in_flight >= 2


def test_budget_never_exceeded_under_workers(store):
    chunks = _chunks(5)

    class Client:
        def __init__(self):
            self.calls = 0
            self.lock = threading.Lock()

        def complete(self, prompt, *, model, params):
            with self.lock:
                self.calls += 1
            return LLMResponse(text="not json", model=model, params=params)

    client = Client()
    drafts, calls, errors = author_drafts(
        chunks,
        client,
        model="m",
        params={"temperature": 0},
        store=store,
        backoff_seconds=0,
        call_budget=3,
        max_workers=4,
    )
    assert client.calls == 3
    assert client.calls <= 3
    assert drafts == []
    assert len(errors) >= 1


def test_one_raising_chunk_isolated_under_workers(store):
    chunks = _chunks(4)

    class Client:
        def complete(self, prompt, *, model, params):
            if "tab2" in prompt:
                raise RuntimeError("boom")
            return LLMResponse(text=_valid_response("F"), model=model, params=params)

    drafts, calls, errors = author_drafts(
        chunks,
        Client(),
        model="m",
        params={"temperature": 0},
        store=store,
        backoff_seconds=0,
        max_workers=4,
    )
    assert len(drafts) == 3
    assert len(errors) == 1
    assert errors[0].chunk_id == "2"
    assert "transport failed" in errors[0].reason


def test_no_leaked_threads_after_workers(store):
    chunks = _chunks(4)

    class Client:
        def complete(self, prompt, *, model, params):
            return LLMResponse(text=_valid_response("F"), model=model, params=params)

    before = set(threading.enumerate())
    author_drafts(
        chunks,
        Client(),
        model="m",
        params={"temperature": 0},
        store=store,
        max_workers=4,
    )
    after = set(threading.enumerate())
    assert after == before


# --- reuse + concurrency -----------------------------------------------------


def _make_mixed_reuse_form(path: Path, tabs: list[str], answers: list[str]) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    for i, tab in enumerate(tabs):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tab
        ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
        ws["A3"] = "Which US states do they ship to"
        ws["F3"] = "X"
        ws["G3"] = "All"
        ws["A4"] = "if not all please check those that apply"
        ws["F4"] = "AK AL AR AZ CA CO CT DE FL GA"
        ws["F5"] = "HI IA ID IL IN KS KY LA MA MD"
        ws["A7"] = "Do they ship to Canada?"
        ws["F7"] = "Yes"
        ws["G7"] = "X"
        ws["H7"] = "No"
        ws["A9"] = "Vendor Status"
        ws["F9"] = "Approved"
        ws["G9"] = "X"
        ws["H9"] = "Provisional"
        ws["J9"] = "Confirm on vendor portal"
        ws["A11"] = "Order Log"
        ws.merge_cells("F11:K11")
        ws["F11"] = "X Support Order Log (Check if the vendor is exporting its order log directly)"
        ws["A13"] = "Order Notes"
        ws["F13"] = answers[i]
        if i % 2 == 1:
            ws["Z20"] = "spacer"
    wb.save(path)
    wb.close()
    return path


def _element_ref(layout, page: int, element_id: str) -> ElementRef:
    for key, bands in layout.bands.items():
        if key // 1000 != page:
            continue
        column = key % 1000
        for band_id, band in enumerate(bands):
            for segment_index, eid in enumerate(band):
                if eid == element_id:
                    region = next(
                        r
                        for r in layout.regions
                        if r.bbox.page == page
                        and r.column == column
                        and band_id in r.band_ids
                    )
                    return ElementRef(
                        region_id=region.region_id,
                        band_id=band_id,
                        segment_index=segment_index,
                    )
    raise AssertionError(f"no layout ref for {element_id!r} on page {page}")


def _reuse_text_responses(path: Path, tabs: list[str]) -> dict[int, str]:
    elements = ingest(path).elements
    layout = analyze(elements)
    responses = {}
    for page in range(len(tabs)):
        label = next(e for e in elements if e.bbox.page == page and e.text == "Order Notes")
        value = next(
            e
            for e in elements
            if e.bbox.page == page
            and e.bbox.y0 == label.bbox.y0
            and e.bbox.x0 > label.bbox.x0
        )
        lref = _element_ref(layout, page, label.element_id)
        vref = _element_ref(layout, page, value.element_id)
        responses[page] = json.dumps(
            {
                "fields": [
                    {
                        "label": "Order Notes",
                        "control_type": "text",
                        "options": [],
                        "answer": [value.text],
                        "annotations": [],
                        "region_id": lref.region_id,
                        "source_elements": [
                            {
                                "region_id": lref.region_id,
                                "band_id": lref.band_id,
                                "segment_index": lref.segment_index,
                            },
                            {
                                "region_id": vref.region_id,
                                "band_id": vref.band_id,
                                "segment_index": vref.segment_index,
                            },
                        ],
                        "confidence": 0.9,
                    }
                ]
            }
        )
    return responses


class _PerTabClient:
    def __init__(self, responses: dict[int, str], tabs: list[str]):
        self.responses = responses
        self.tabs = tabs
        self.calls = 0

    def complete(self, prompt, *, model, params):
        self.calls += 1
        tab = re.search(r"Layout projection for `([^`]+)`", prompt).group(1)
        page = self.tabs.index(tab)
        return LLMResponse(text=self.responses[page], model=model, params=params)


def test_reuse_serial_equals_workers(tmp_path: Path):
    tabs = [f"Tab {i}" for i in range(4)]
    path = _make_mixed_reuse_form(
        tmp_path / "mixed.xlsx", tabs, ["a", "b", "c", "d"]
    )
    responses = _reuse_text_responses(path, tabs)

    def run(workers: int):
        client = _PerTabClient(responses, tabs)
        record = Pipeline(
            Store(tmp_path / f"store-{workers}"),
            client,
            PipelineConfig(reuse_layout_bindings=True, chunk_workers=workers),
        ).run(path)
        return record, client

    rec1, client1 = run(1)
    rec4, client4 = run(4)

    assert client1.calls == client4.calls
    assert rec1.status is InstanceStatus.COMPLETE
    assert rec4.status is InstanceStatus.COMPLETE

    def comparable(record):
        return [
            (
                f.tab,
                f.label_text,
                f.field_id,
                f.value_normalized,
                tuple(f.source_elements),
                f.provenance.source.value if f.provenance.source else None,
            )
            for f in record.fields
        ]

    assert comparable(rec1) == comparable(rec4)
    assert sum(1 for f in rec4.fields if f.provenance.source is ProvenanceSource.REPLAY) == 2


def test_chunk_workers_validated(tmp_path):
    store = Store(tmp_path / "store")
    for bad in (0, 33, True, False, "4", 1.5):
        with pytest.raises(ValueError):
            Pipeline(store, None, PipelineConfig(chunk_workers=bad))
    Pipeline(store, None, PipelineConfig(chunk_workers=32))

