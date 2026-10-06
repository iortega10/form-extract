from __future__ import annotations

import json
import re
from pathlib import Path

from formextract.model import InstanceStatus, ProvenanceSource
from formextract.pipeline import Pipeline, PipelineConfig, compute_cache_key
from formextract.resolve import LLMResponse
from formextract.store import Store


class _TabClient:
    def __init__(self):
        self.calls = 0

    def complete(self, prompt: str, *, model: str, params: dict) -> LLMResponse:
        self.calls += 1
        match = re.search(r"Layout projection for `([^`]+)`", prompt)
        tab = match.group(1) if match else "?"
        return LLMResponse(
            text=json.dumps(
                {
                    "fields": [
                        {
                            "label": f"Field on {tab}",
                            "control_type": "text",
                            "answer": [tab],
                            "annotations": [],
                        }
                    ]
                }
            ),
            model=model,
            params=params,
            tokens=1,
        )


def _make_multi_tab(path: Path, n: int) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    for i in range(n):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = f"Checklist {i}"
        ws["A1"] = f"Field on Checklist {i}"
    wb.save(path)
    wb.close()
    return path


def test_flag_off_matches_the_existing_path(tmp_path: Path):
    path = _make_multi_tab(tmp_path / "multi.xlsx", n=6)
    client = _TabClient()
    record = Pipeline(Store(tmp_path / "store"), client, PipelineConfig()).run(path)

    assert record.status is InstanceStatus.COMPLETE
    assert client.calls == 6  # one authoring call per chunk, no reuse
    assert len(record.fields) == 6
    assert all(f.provenance.source is ProvenanceSource.LLM for f in record.fields)
    assert {f.label_text for f in record.fields} == {
        f"Field on Checklist {i}" for i in range(6)
    }


def test_cache_key_varies_with_reuse_flag():
    base = dict(
        content_hash="abc123",
        pipeline_version="2",
        schema_version="2",
        prompt_version="2",
        model="canned",
        params={"temperature": 0},
    )
    off = compute_cache_key(reuse_layout_bindings=False, **base)
    on = compute_cache_key(reuse_layout_bindings=True, **base)
    assert off != on
