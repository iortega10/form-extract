"""Field.tab for a region-less draft (0.6.0-G2 part 4).

A projection chunk always knows its tab, so a draft whose ``region_id`` does
not resolve must still name that tab. ``author_drafts`` carries the chunk key
onto the draft and ``drafts_to_fields`` uses it when the region is missing; the
field_id (page/column/region/anchor/ordinal) never sees the tab, so it does not
move.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import (
    AuthoredBy,
    BindingDraft,
    BindingProvenance,
    BBox,
    ControlType,
    ElementRef,
    Option,
)
from formextract.pipeline import (
    Pipeline,
    PipelineConfig,
    _normalize_checkbox_conventions,
)
from formextract.resolve import (
    DECLARED_CONVENTION_TOKEN,
    LLMResponse,
    _field_id,
    drafts_to_fields,
    remap_drafts_for_page,
)
from formextract.schema import normalize_label
from formextract.store import Store

LABEL = "Alpha widget variant?"
CELLS = (1, 3, 4, 5, 6)  # label, marker, option, marker, option


def _write_tabs(path: Path, titles: list[str]) -> None:
    wb = Workbook()
    for i, title in enumerate(titles):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws["A1"] = LABEL
        ws["C1"] = "X"
        ws["D1"] = "Ruby"
        ws["E1"] = "X"
        ws["F1"] = "Teal"
    wb.save(path)
    wb.close()


def _ref_for(layout, element_id: str) -> ElementRef:
    for key, bands in layout.bands.items():
        column = key % 1000
        for band_id, band in enumerate(bands):
            for segment, eid in enumerate(band):
                if eid == element_id:
                    region = next(
                        r
                        for r in layout.regions
                        if r.column == column and band_id in r.band_ids
                    )
                    return ElementRef(
                        region_id=region.region_id,
                        band_id=band_id,
                        segment_index=segment,
                    )
    raise AssertionError(element_id)


def _layout_for(path: Path, titles: list[str]):
    ing = ingest(path)
    layout = analyze(ing.elements)
    by_id = {e.element_id: e for e in ing.elements}
    refs = {t: [_ref_for(layout, f"{t}!1:{c}") for c in CELLS] for t in titles}
    return layout, by_id, refs


def _draft(*, tab, region_id, refs) -> BindingDraft:
    return BindingDraft(
        draft_id="d-1",
        label=LABEL,
        control_type=ControlType.SINGLE_SELECT,
        bbox=BBox(),
        options=[Option(text="Ruby", selected=True), Option(text="Teal", selected=False)],
        region_id=region_id,
        source_refs=list(refs),
        provenance=BindingProvenance(authored_by=AuthoredBy.LLM, authored_at="now"),
        tab=tab,
    )


def _response(*, region_id: str, refs) -> str:
    return json.dumps(
        {
            "fields": [
                {
                    "label": LABEL,
                    "control_type": "single_select",
                    "options": [
                        {"text": "Ruby", "selected": True},
                        {"text": "Teal", "selected": False},
                    ],
                    "answer": [],
                    "annotations": [],
                    "region_id": region_id,
                    "source_elements": [
                        {
                            "region_id": r.region_id,
                            "band_id": r.band_id,
                            "segment_index": r.segment_index,
                        }
                        for r in refs
                    ],
                }
            ]
        }
    )


class _Client:
    def __init__(self, text: str):
        self.text = text
        self.calls = 0

    def complete(self, prompt: str, *, model: str, params: dict) -> LLMResponse:
        self.calls += 1
        return LLMResponse(
            text=self.text, model=model, params=params, tokens=1, latency_ms=0
        )


def _run_pipeline(tmp_path, region_id: str):
    path = tmp_path / "wb.xlsx"
    _write_tabs(path, ["TheTab"])
    _layout, _by_id, refs = _layout_for(path, ["TheTab"])
    client = _Client(_response(region_id=region_id, refs=refs["TheTab"]))
    record = Pipeline(Store(tmp_path / "store"), client, PipelineConfig()).run(path)
    assert len(record.fields) == 1
    return record.fields[0]


def test_region_less_field_carries_the_chunk_tab(tmp_path):
    field = _run_pipeline(tmp_path, region_id="no-such-region")
    # the region did not resolve...
    assert field.region_ref == "no-such-region"
    # ... but the chunk's tab still names the field
    assert field.tab == "TheTab"
    assert field.source_elements


def test_region_resolved_field_tab_matches_the_chunk(tmp_path):
    path = tmp_path / "wb.xlsx"
    _write_tabs(path, ["TheTab"])
    _layout, _by_id, refs = _layout_for(path, ["TheTab"])
    field = _run_pipeline(tmp_path, region_id=refs["TheTab"][0].region_id)
    assert field.region_ref == refs["TheTab"][0].region_id
    # the region-derived tab and the carried chunk tab agree
    assert field.tab == "TheTab"


def _conventions(*selectors):
    return _normalize_checkbox_conventions(list(selectors))


def _fields_for(tmp_path, draft, conventions):
    path = tmp_path / "wb.xlsx"
    _write_tabs(path, ["TheTab"])
    layout, by_id, _refs = _layout_for(path, ["TheTab"])
    return drafts_to_fields(
        [draft],
        layout=layout,
        elements_by_id=by_id,
        tabs=["TheTab"],
        checkbox_conventions=conventions,
        page_tabs={0: "TheTab"},
    )


def test_declared_convention_applies_to_a_region_less_field(tmp_path):
    path = tmp_path / "wb.xlsx"
    _write_tabs(path, ["TheTab"])
    _layout, _by_id, refs = _layout_for(path, ["TheTab"])
    draft = _draft(tab="TheTab", region_id="no-such-region", refs=refs["TheTab"])
    fields = _fields_for(
        tmp_path,
        draft,
        _conventions({"tab": "TheTab", "convention": "mark_precedes_option"}),
    )
    field = fields[0]
    assert field.tab == "TheTab"
    # the between-marker selects the option on its right, via the declaration
    assert [o.text for o in field.options if o.selected] == ["Teal"]
    assert DECLARED_CONVENTION_TOKEN in field.provenance.heuristic_agreement


def test_the_tab_decides_which_convention_applies(tmp_path):
    """Before part 4 a region-less draft matched *every* selector.

    ``_matching_conventions`` skips the tab filter when the tab is None, so the
    longest / most specific selector won even when it belonged to another tab:
    here ``OtherTab`` (mark_follows_option -> "Ruby") beat ``TheTab``
    (mark_precedes_option -> "Teal"). With the carried tab the right one wins.
    """
    path = tmp_path / "wb.xlsx"
    _write_tabs(path, ["TheTab"])
    _layout, _by_id, refs = _layout_for(path, ["TheTab"])
    conventions = _conventions(
        {"tab": "OtherTab", "convention": "mark_follows_option"},
        {"tab": "TheTab", "convention": "mark_precedes_option"},
    )
    with_tab = _fields_for(
        tmp_path,
        _draft(tab="TheTab", region_id="no-such-region", refs=refs["TheTab"]),
        conventions,
    )[0]
    without_tab = _fields_for(
        tmp_path,
        _draft(tab=None, region_id="no-such-region", refs=refs["TheTab"]),
        conventions,
    )[0]

    assert [o.text for o in with_tab.options if o.selected] == ["Teal"]
    assert [o.text for o in without_tab.options if o.selected] == ["Ruby"]
    assert with_tab.tab == "TheTab"
    assert without_tab.tab is None


def test_field_id_is_identical_with_and_without_a_tab(tmp_path):
    path = tmp_path / "wb.xlsx"
    _write_tabs(path, ["TheTab"])
    _layout, _by_id, refs = _layout_for(path, ["TheTab"])
    without_tab = _fields_for(
        tmp_path,
        _draft(tab=None, region_id="no-such-region", refs=refs["TheTab"]),
        [],
    )[0]
    with_tab = _fields_for(
        tmp_path,
        _draft(tab="TheTab", region_id="no-such-region", refs=refs["TheTab"]),
        [],
    )[0]

    assert with_tab.field_id == without_tab.field_id
    # the region-less field_id is computed with page=0, column=-1, the draft's
    # region id and the label's ordinal only: the tab never enters the key.
    assert with_tab.field_id == _field_id(
        0, -1, "no-such-region", normalize_label(LABEL), 0
    )


def test_replay_path_names_the_target_tab(tmp_path):
    path = tmp_path / "wb.xlsx"
    _write_tabs(path, ["S1", "S2"])
    layout, by_id, refs = _layout_for(path, ["S1", "S2"])
    exemplar = _draft(tab="S1", region_id=refs["S1"][0].region_id, refs=refs["S1"])
    region_bbox = {r.region_id: r.bbox for r in layout.regions}
    remapped = remap_drafts_for_page(
        [exemplar], 0, 1, regions=layout.regions, region_bbox=region_bbox
    )
    fields = drafts_to_fields(
        remapped,
        layout=layout,
        elements_by_id=by_id,
        tabs=["S1", "S2"],
        page_tabs={0: "S1", 1: "S2"},
    )
    assert fields
    # a replayed draft always resolves its region, so the target tab comes from
    # the page map and never from the exemplar
    assert {f.tab for f in fields} == {"S2"}
