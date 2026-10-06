from __future__ import annotations

import json

from formextract.evals.structure_probe import ALLOWED_KEYS, run_probe


def _make_workbook(path, rows, title="Checklist", extra_rows=None):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r, (label, cells) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=label)
        for col, text in cells:
            ws[f"{col}{r}"] = text
    for r, (label, cells) in enumerate(extra_rows or [], start=len(rows) + 1):
        ws.cell(row=r, column=1, value=label)
        for col, text in cells:
            ws[f"{col}{r}"] = text
    wb.save(path)
    wb.close()


def _between_rows():
    return [("Priority", [("F", "Low"), ("G", "X"), ("H", "High")])]


def test_probe_output_keys_are_exactly_allowed(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, _between_rows())
    result = run_probe(path)
    assert set(result) == ALLOWED_KEYS


def test_probe_values_are_ints_and_one_bool(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, _between_rows())
    result = run_probe(path)
    for key, value in result.items():
        if key == "markers_total_equals_bucket_sum":
            assert isinstance(value, bool)
        elif key == "region_counts_by_tab":
            assert isinstance(value, list)
            assert all(isinstance(v, int) for v in value)
        else:
            assert isinstance(value, int), key


def test_probe_does_not_leak_text(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(
        path,
        _between_rows(),
        extra_rows=[("ZZTOP-SECRET-42", [("F", "Distinctive Vendor Name")])],
    )
    result = run_probe(path)
    text = json.dumps(result, sort_keys=True)
    assert "ZZTOP-SECRET-42" not in text
    assert "Distinctive Vendor Name" not in text
    assert "Priority" not in text
    assert "Low" not in text


def test_probe_bucket_sum_matches_marker_total(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, _between_rows())
    result = run_probe(path)
    assert result["markers_between"] == 1
    assert result["markers_total_equals_bucket_sum"] is True


def test_probe_reports_convention_selection(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, _between_rows())
    conventions = [
        {
            "tab": "Checklist",
            "anchor_pattern": "priority",
            "convention": "mark_follows_option",
        }
    ]
    result = run_probe(path, conventions)
    assert result["controls_convention_selected"] == 1
    assert result["controls_ambiguous"] == 0
