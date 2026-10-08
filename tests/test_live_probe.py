from __future__ import annotations

import contextlib
import http.server
import json
import sys
import threading
import time
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
if str(REPO / "tools") not in sys.path:
    sys.path.insert(0, str(REPO / "tools"))

import live_probe as probe  # noqa: E402
import make_gold  # noqa: E402

FIELD_LABEL = "ZZTOP-SECRET-LABEL"
OK_FIELDS = (
    '{"fields": [{"label": "' + FIELD_LABEL + '", "control_type": "text", '
    '"options": [], "answer": ["v"], "annotations": []}]}'
)
TRUNCATED_FIELDS = '{"fields": [{"label": "' + FIELD_LABEL + '", "control_type": "text"'


def _openai_body(text: str, finish: str) -> bytes:
    return json.dumps(
        {
            "choices": [{"message": {"content": text}, "finish_reason": finish}],
            "usage": {
                "prompt_tokens": 11,
                "completion_tokens": 7,
                "completion_tokens_details": {"reasoning_tokens": 3},
            },
        }
    ).encode()


def _gemini_body(text: str, finish: str) -> bytes:
    return json.dumps(
        {
            "candidates": [
                {"content": {"parts": [{"text": text}]}, "finishReason": finish}
            ],
            "usageMetadata": {
                "promptTokenCount": 11,
                "candidatesTokenCount": 7,
                "thoughtsTokenCount": 3,
            },
        }
    ).encode()


@contextlib.contextmanager
def _fake_server(*, ok_body: bytes, trunc_body: bytes, key: str):
    state = {
        "requests": [],
        "counts": {},
        "key": key,
        "ok_body": ok_body,
        "trunc_body": trunc_body,
    }

    class Handler(http.server.BaseHTTPRequestHandler):
        def log_message(self, *args):  # noqa: D102 - silence the test server
            pass

        def do_POST(self):  # noqa: N802 - http.server API
            length = int(self.headers.get("Content-Length", 0))
            self.rfile.read(length)
            state["requests"].append(
                {
                    "path": self.path,
                    "requestline": self.requestline,
                    "headers": {k.lower(): v for k, v in self.headers.items()},
                }
            )
            route = self.path.split("?", 1)[0].strip("/").split("/")[0]
            state["counts"][route] = state["counts"].get(route, 0) + 1
            if route == "e401":
                payload = (
                    '{"error": {"message": "invalid key: ' + key + '"}}'
                ).encode()
                self._send(401, payload)
            elif route == "e400":
                self._send(400, b'{"error": {"message": "bad request"}}')
            elif route == "e429":
                self._send(429, b'{"error": {"message": "slow down"}}')
            elif route == "slow":
                time.sleep(0.3)
                self._send(200, ok_body)
            elif route == "trunc":
                self._send(200, trunc_body)
            else:
                self._send(200, ok_body)

        def _send(self, code: int, payload: bytes) -> None:
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    server.daemon_threads = True
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield server, f"http://127.0.0.1:{server.server_address[1]}", state
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def _wait_idle(before: int) -> int:
    deadline = time.time() + 3
    while threading.active_count() != before and time.time() < deadline:
        time.sleep(0.02)
    return threading.active_count()


def _run(argv: list[str]) -> int:
    try:
        return probe.main(argv)
    except SystemExit as exc:
        return exc.code if isinstance(exc.code, int) else 1


def _make_workbook(path: Path, sheets: tuple[str, ...] = ("Alpha Log",)) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    first = wb.active
    first.title = sheets[0]
    for name in sheets:
        ws = first if name == sheets[0] else wb.create_sheet(name)
        ws["A1"] = name
        ws["A2"] = "Priority"
        ws["B2"] = "Low"
        ws["C2"] = "X"
        ws["D2"] = "High"
    wb.save(path)
    wb.close()
    return path


def _assert_scalar(value) -> None:
    if value is None or isinstance(value, (int, float, bool)):
        return
    assert isinstance(value, str) and value in probe.ALLOWED_STRING_VALUES, value


def _last_line(out: str) -> dict:
    lines = [line for line in out.splitlines() if line.strip()]
    assert lines
    return json.loads(lines[-1])


def test_stdout_is_numbers_only_and_leaks_no_text(tmp_path, capsys):
    path = _make_workbook(tmp_path / "wb.xlsx")
    before = threading.active_count()
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=""
    ) as (_server, base, _state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
            ]
        )
        out, err = capsys.readouterr()

    assert err == ""
    assert code == 0
    lines = [line for line in out.splitlines() if line.strip()]
    assert len(lines) == 1
    data = json.loads(lines[0])
    for value in data.values():
        _assert_scalar(value)
    assert FIELD_LABEL not in out
    assert "control_type" not in out
    # row coverage (T4): the denominator is anchors + GRID-region units, and
    # the old ad-hoc key is gone
    assert "anchors_consumed" not in data
    assert data["units_total"] == data["anchors"] + data["grid_units"]
    assert 0 <= data["units_consumed"] <= data["units_total"]
    assert "annotations" not in out
    assert data["chunks_run"] == 1
    assert data["fields"] == 1
    assert data["prompt_tokens"] == 11
    assert data["output_tokens"] == 7
    assert data["reasoning_tokens"] == 3
    assert data["status_class"] == "complete"
    assert data["http_error_class"] == "none"
    assert _wait_idle(before) == before


def test_key_never_appears_even_when_the_body_echoes_it(tmp_path, capsys, monkeypatch):
    key = "test-key-not-real"
    monkeypatch.setenv("PROBE_TEST_KEY", key)
    path = _make_workbook(tmp_path / "wb.xlsx")
    before = threading.active_count()
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=key
    ) as (_server, base, _state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/e401",
                "--env-var", "PROBE_TEST_KEY",
            ]
        )
        out, err = capsys.readouterr()

    assert code == 0
    assert key not in out
    assert key not in err
    data = _last_line(out)
    assert data["http_error_class"] == "4xx"
    assert _wait_idle(before) == before


def test_redaction_scrubs_message_and_traceback():
    key = "test-key-not-real"
    secrets = {"test-key-not-real", "Bearer test-key-not-real"}
    exc = RuntimeError(f"call failed with {key}")
    name, message, traceback_text = probe.render_error(exc, secrets)
    assert name == "RuntimeError"
    assert key not in message
    assert key not in traceback_text
    assert "<redacted>" in message


def test_unexpected_error_is_reported_redacted(tmp_path, capsys, monkeypatch):
    key = "test-key-not-real"
    monkeypatch.setenv("PROBE_TEST_KEY", key)
    path = _make_workbook(tmp_path / "wb.xlsx")

    def boom(*args, **kwargs):
        raise RuntimeError(f"install failed for {key}")

    monkeypatch.setattr(probe, "ingest", boom)
    code = _run(
        [
            str(path),
            "--provider", "openai",
            "--model", "test-model",
            "--base-url", "http://127.0.0.1:1/ok",
            "--env-var", "PROBE_TEST_KEY",
        ]
    )
    out, err = capsys.readouterr()

    assert code == 1
    assert key not in out
    assert key not in err
    assert "RuntimeError" in err


def test_4xx_is_classified_and_costs_one_request(tmp_path, capsys):
    path = _make_workbook(tmp_path / "wb.xlsx")
    before = threading.active_count()
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=""
    ) as (_server, base, state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/e400",
                "--env-var", "",
            ]
        )
        out, _err = capsys.readouterr()

    assert code == 0
    assert state["counts"].get("e400") == 1
    data = _last_line(out)
    assert data["http_error_class"] == "4xx"
    assert data["status_class"] == "failed"
    assert _wait_idle(before) == before


def test_length_cap_and_truncated_are_reported(tmp_path, capsys):
    path = _make_workbook(tmp_path / "wb.xlsx")
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"),
        trunc_body=_openai_body(TRUNCATED_FIELDS, "length"),
        key="",
    ) as (_server, base, _state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/trunc",
                "--env-var", "",
            ]
        )
        out, _err = capsys.readouterr()

    assert code == 0
    data = _last_line(out)
    assert data["hit_length_cap"] is True
    assert data["truncated"] is True
    assert data["finish_reason_class"] == "length"
    assert data["fields"] == 0
    assert data["status_class"] == "partial"


def test_gemini_key_is_only_in_the_header(tmp_path, capsys, monkeypatch):
    key = "test-key-not-real"
    monkeypatch.setenv("PROBE_TEST_KEY", key)
    path = _make_workbook(tmp_path / "wb.xlsx")
    before = threading.active_count()
    with _fake_server(
        ok_body=_gemini_body(OK_FIELDS, "STOP"), trunc_body=b"", key=key
    ) as (_server, base, state):
        code = _run(
            [
                str(path),
                "--provider", "gemini",
                "--model", "gemini-test",
                "--base-url", base + "/ok",
                "--env-var", "PROBE_TEST_KEY",
            ]
        )
        out, _err = capsys.readouterr()

    assert code == 0
    assert state["requests"]
    for request in state["requests"]:
        assert key not in request["path"]
        assert key not in request["requestline"]
        assert "key=" not in request["path"]
        assert request["headers"].get("x-goog-api-key") == key
        assert "authorization" not in request["headers"]
    assert _last_line(out)["fields"] == 1
    assert _wait_idle(before) == before


def test_openai_key_is_in_the_authorization_header(tmp_path, capsys, monkeypatch):
    key = "test-key-not-real"
    monkeypatch.setenv("PROBE_TEST_KEY", key)
    path = _make_workbook(tmp_path / "wb.xlsx")
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=key
    ) as (_server, base, state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "PROBE_TEST_KEY",
            ]
        )
        capsys.readouterr()

    assert code == 0
    for request in state["requests"]:
        assert key not in request["path"]
        assert request["headers"].get("authorization") == "Bearer " + key


def test_missing_env_var_exits_2(tmp_path, capsys, monkeypatch):
    monkeypatch.delenv("PROBE_UNSET_KEY_XYZ", raising=False)
    path = _make_workbook(tmp_path / "wb.xlsx")
    code = _run(
        [
            str(path),
            "--provider", "openai",
            "--model", "test-model",
            "--env-var", "PROBE_UNSET_KEY_XYZ",
        ]
    )
    out, err = capsys.readouterr()

    assert code == 2
    assert out == ""
    assert "key variable PROBE_UNSET_KEY_XYZ is not set" in err


def test_tab_takes_an_index_and_rejects_a_name(tmp_path, capsys):
    path = _make_workbook(tmp_path / "wb.xlsx", sheets=("Alpha Log", "Beta Log"))
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=""
    ) as (_server, base, _state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
                "--tab", "1",
            ]
        )
        out, _err = capsys.readouterr()

    assert code == 0
    assert _last_line(out)["chunks_run"] == 1

    code_name = _run(
        [
            str(path),
            "--provider", "openai",
            "--model", "test-model",
            "--base-url", base + "/ok",
            "--env-var", "",
            "--tab", "Alpha Log",
        ]
    )
    assert code_name == 2

    code_range = _run(
        [
            str(path),
            "--provider", "openai",
            "--model", "test-model",
            "--base-url", base + "/ok",
            "--env-var", "",
            "--tab", "9",
        ]
    )
    assert code_range == 1


def test_workers_match_serial_counts(tmp_path, capsys):
    path = _make_workbook(tmp_path / "wb.xlsx", sheets=("Alpha Log", "Beta Log"))
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=""
    ) as (_server, base, _state):
        code1 = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
                "--workers", "1",
            ]
        )
        out1, _err1 = capsys.readouterr()
        code2 = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
                "--workers", "4",
            ]
        )
        out2, _err2 = capsys.readouterr()

    assert code1 == 0 and code2 == 0
    serial = json.loads(out1.strip())
    pooled = json.loads(out2.strip())
    assert serial["chunks_run"] == 2 and pooled["chunks_run"] == 2
    for key, value in serial.items():
        if key in ("seconds", "tokens_per_second"):
            continue
        assert pooled[key] == value, key


def test_tiny_reports_floor_seconds(tmp_path, capsys):
    with _fake_server(
        ok_body=_openai_body('{"ok": true}', "stop"), trunc_body=b"", key=""
    ) as (_server, base, _state):
        code = _run(
            [
                "--tiny",
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
            ]
        )
        out, _err = capsys.readouterr()

    assert code == 0
    data = _last_line(out)
    assert isinstance(data["floor_seconds"], float)
    assert data["output_tokens"] == 7


def test_path_is_required_without_tiny():
    assert _run(["--provider", "openai", "--model", "test-model"]) == 2


def test_repeat_prints_one_aggregated_object(tmp_path, capsys):
    path = _make_workbook(tmp_path / "wb.xlsx")
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=""
    ) as (_server, base, _state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
                "--repeat", "3",
            ]
        )
        out, _err = capsys.readouterr()

    assert code == 0
    lines = [line for line in out.splitlines() if line.strip()]
    assert len(lines) == 1
    data = json.loads(lines[0])
    for value in data.values():
        _assert_scalar(value)
    assert data["repeat"] == 3
    assert data["fields_min"] == 1
    assert data["fields_max"] == 1
    assert data["truncation_count"] == 0


def test_dump_writes_text_only_when_asked(tmp_path, capsys):
    path = _make_workbook(tmp_path / "wb.xlsx")
    dump = tmp_path / "dump.json"
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=""
    ) as (_server, base, _state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
                "--dump", str(dump),
            ]
        )
        out, err = capsys.readouterr()

    assert code == 0
    assert FIELD_LABEL not in out
    assert "may contain private data" in err
    payload = json.loads(dump.read_text(encoding="utf-8"))
    assert payload["calls"][0]["prompt"]


# --------------------------------------------------------------------------
# --gold mode: run every gold tab through the real Pipeline and score it
# --------------------------------------------------------------------------

_CONTROL = {
    "single": "single_select",
    "multi": "multi_select",
    "bool": "bool",
    "text": "text",
}


def _ref_for(layout, element_id: str):
    from formextract.model import ElementRef

    for key, bands in layout.bands.items():
        column = key % 1000
        for band_id, band in enumerate(bands):
            for segment, eid in enumerate(band):
                if eid == element_id:
                    region = next(
                        r
                        for r in layout.regions
                        if r.column == column and band_id in r.band_ids
                    )
                    return ElementRef(
                        region_id=region.region_id,
                        band_id=band_id,
                        segment_index=segment,
                    )
    raise AssertionError(element_id)


def _ref_dict(ref) -> dict:
    return {
        "region_id": ref.region_id,
        "band_id": ref.band_id,
        "segment_index": ref.segment_index,
    }


def _fields_for_tab(gold_doc: dict, tab: str, layout, *, swap_single=False) -> list:
    cells = {c["id"]: c for c in gold_doc["cells"]}
    out = []
    for field in gold_doc["fields"]:
        if field["tab"] != tab:
            continue
        cell_ids = (
            field["label_cells"]
            + field["option_cells"]
            + field["answer_cells"]
            + field["annotation_cells"]
        )
        refs = {cid: _ref_dict(_ref_for(layout, cid)) for cid in cell_ids}
        control = _CONTROL[field["kind"]]
        if swap_single and control == "single_select":
            control = "multi_select"
        out.append(
            {
                "label": " ".join(cells[c]["text"] for c in field["label_cells"]),
                "control_type": control,
                "options": [
                    {
                        "text": cells[cid]["text"],
                        "selected": cid in field["selected_option_cells"],
                    }
                    for cid in field["option_cells"]
                ],
                "answer": (
                    [field["answer_text"]]
                    if field["kind"] == "text" and field["answer_text"]
                    else []
                ),
                "annotations": [
                    cells[c]["text"] for c in field["annotation_cells"]
                ],
                "region_id": refs[cell_ids[0]]["region_id"],
                "source_elements": [refs[c] for c in cell_ids],
            }
        )
    return out


def _gold_payloads(gold_dir, gold_doc, *, keep=None, swap_single=False) -> list:
    """One OpenAI-style response body per tab, built from the real layout."""
    from formextract.ingest import ingest
    from formextract.layout import analyze

    bodies = []
    for tab in gold_doc["tabs"]:
        ing = ingest(gold_dir / f"{tab}.xlsx")
        layout = analyze(ing.elements)
        fields = _fields_for_tab(gold_doc, tab, layout, swap_single=swap_single)
        if keep is not None:
            fields = [f for i, f in enumerate(fields) if keep(i)]
        bodies.append(_openai_body(json.dumps({"fields": fields}), "stop"))
    return bodies


@contextlib.contextmanager
def _scripted_server(responses: list, *, key: str = ""):
    state = {"requests": [], "counts": {}, "index": 0, "key": key, "responses": responses}

    class Handler(http.server.BaseHTTPRequestHandler):
        def log_message(self, *args):  # noqa: D102 - silence the test server
            pass

        def do_POST(self):  # noqa: N802 - http.server API
            length = int(self.headers.get("Content-Length", 0))
            self.rfile.read(length)
            state["requests"].append(
                {
                    "path": self.path,
                    "headers": {k.lower(): v for k, v in self.headers.items()},
                }
            )
            route = self.path.split("?", 1)[0].strip("/").split("/")[0]
            state["counts"][route] = state["counts"].get(route, 0) + 1
            if route == "e401":
                payload = (
                    '{"error": {"message": "invalid key: ' + key + '"}}'
                ).encode()
                self._send(401, payload)
                return
            i = state["index"]
            state["index"] = i + 1
            self._send(200, state["responses"][min(i, len(state["responses"]) - 1)])

        def _send(self, code: int, payload: bytes) -> None:
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    server.daemon_threads = True
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield server, f"http://127.0.0.1:{server.server_address[1]}", state
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


@pytest.fixture(scope="module")
def gold_dir(tmp_path_factory):
    base = tmp_path_factory.mktemp("probe-gold")
    original = make_gold.HELDOUT_TABS
    make_gold.HELDOUT_TABS = (
        ("PYesNo", make_gold._tab_yes_no),
        ("PDense", make_gold._tab_dense),
        ("PText", make_gold._tab_text_and_typed),
        ("PSide", make_gold._tab_side_by_side),
    )
    try:
        gold, _manifest, _ = make_gold.build_set(
            "heldout", make_gold.DEFAULT_SEEDS["heldout"], base
        )
    finally:
        make_gold.HELDOUT_TABS = original
    return base, gold


def _gold_args(base: str, gold_path, *extra: str) -> list:
    return [
        "--gold", str(gold_path),
        "--provider", "openai",
        "--model", "test-model",
        "--base-url", base,
        "--env-var", "",
        *extra,
    ]


def test_gold_mode_scores_a_perfect_server_run(gold_dir, capsys):
    gold_path, gold_doc = gold_dir
    bodies = _gold_payloads(gold_path, gold_doc)
    before = threading.active_count()
    with _scripted_server(bodies) as (_server, base, state):
        code = _run(_gold_args(base, gold_path))
        out, err = capsys.readouterr()

    assert code == 0
    assert err == ""
    data = _last_line(out)
    assert data["gold_fields"] == len(gold_doc["fields"])
    assert data["matched"] == data["gold_fields"]
    assert data["matched_ignoring_kind"] == data["gold_fields"]
    assert data["kind_confusions"] == 0
    assert data["kind_confusion_pairs"] == {}
    assert data["recall"] == 1.0
    assert data["precision"] == 1.0
    assert data["recall_num"] == data["recall_den"] == data["gold_fields"]
    assert data["missed"] == 0 and data["spurious"] == 0
    assert data["unexpected_ambiguous"] == 0
    assert data["selected_ok_under_convention"] + data["selected_ambiguous_expected"] == data["selected_ok"]
    assert data["calls"] == len(gold_doc["tabs"])
    assert data["per_tag"]
    assert sum(t["gold"] for t in data["per_tag"].values()) == data["gold_fields"]
    assert state["index"] == len(gold_doc["tabs"])
    assert _wait_idle(before) == before


def test_gold_mode_forwards_the_kind_counters(gold_dir, capsys):
    """The probe forwards the grouping/kind split, overall and per tag."""
    gold_path, gold_doc = gold_dir
    # every single-select field read as a multi-select: grouping intact, only the
    # kind counters move.
    bodies = _gold_payloads(gold_path, gold_doc, swap_single=True)
    with _scripted_server(bodies) as (_server, base, _state):
        code = _run(_gold_args(base, gold_path))
        out, _err = capsys.readouterr()

    assert code == 0
    data = _last_line(out)
    assert data["gold_fields"] == len(gold_doc["fields"])
    assert data["matched_ignoring_kind"] == data["gold_fields"]
    assert data["kind_confusions"] > 0
    assert data["matched"] == data["matched_ignoring_kind"] - data["kind_confusions"]
    assert data["kind_confusion_pairs"] == {"single>multi": data["kind_confusions"]}
    # per tag: the split travels too, and sums back to the overall
    assert data["per_tag"]
    assert sum(t["matched_ignoring_kind"] for t in data["per_tag"].values()) > 0
    assert all("kind_confusions" in t for t in data["per_tag"].values())
    n_single = sum(
        1 for f in gold_doc["fields"] if f["kind"] == "single"
    )
    assert data["kind_confusions"] == n_single


def test_gold_mode_mutation_lowers_recall(gold_dir, capsys):
    gold_path, gold_doc = gold_dir
    kept = sum(
        len(range(0, sum(1 for f in gold_doc["fields"] if f["tab"] == tab), 2))
        for tab in gold_doc["tabs"]
    )
    bodies = _gold_payloads(gold_path, gold_doc, keep=lambda i: i % 2 == 0)
    with _scripted_server(bodies) as (_server, base, _state):
        code = _run(_gold_args(base, gold_path))
        out, _err = capsys.readouterr()

    total = len(gold_doc["fields"])
    assert code == 0
    data = _last_line(out)
    assert data["matched"] == kept
    assert data["missed"] == total - kept
    assert data["recall_num"] == kept
    assert data["recall_den"] == total
    assert data["recall"] == round(kept / total, 6)


def test_gold_mode_output_is_numbers_only(gold_dir, capsys):
    gold_path, gold_doc = gold_dir
    bodies = _gold_payloads(gold_path, gold_doc)
    with _scripted_server(bodies) as (_server, base, _state):
        code = _run(_gold_args(base, gold_path))
        out, _err = capsys.readouterr()

    assert code == 0
    lines = [line for line in out.splitlines() if line.strip()]
    assert len(lines) == 1
    data = json.loads(lines[0])
    for key, value in data.items():
        if key in (
            "per_tag",
            "finish_reason_class_counts",
            "kind_confusion_pairs",
            "http_status_counts",
        ):
            for inner in value.values():
                assert isinstance(inner, (int, dict))
            continue
        _assert_scalar(value)
    # no gold cell text or tab name survives into the output
    secrets = {c["text"] for c in gold_doc["cells"]} | set(gold_doc["tabs"])
    for secret in secrets:
        assert secret not in out


def test_gold_mode_key_never_appears_in_output(gold_dir, capsys, monkeypatch):
    gold_path, gold_doc = gold_dir
    key = "test-key-not-real"
    monkeypatch.setenv("PROBE_TEST_KEY", key)
    bodies = _gold_payloads(gold_path, gold_doc)
    args = _gold_args("", gold_path)
    args[args.index("--env-var") + 1] = "PROBE_TEST_KEY"
    with _scripted_server(bodies, key=key) as (_server, base, _state):
        args[args.index("--base-url") + 1] = base
        code = _run(args)
        out, err = capsys.readouterr()

    assert code == 0
    assert key not in out
    assert key not in err


def test_gold_mode_401_body_echo_is_redacted(gold_dir, capsys, monkeypatch):
    gold_path, gold_doc = gold_dir
    key = "test-key-not-real"
    monkeypatch.setenv("PROBE_TEST_KEY", key)
    bodies = _gold_payloads(gold_path, gold_doc)
    args = _gold_args("", gold_path)
    args[args.index("--env-var") + 1] = "PROBE_TEST_KEY"
    with _scripted_server(bodies, key=key) as (_server, base, state):
        args[args.index("--base-url") + 1] = base + "/e401"
        code = _run(args)
        out, err = capsys.readouterr()

    assert code == 0
    assert key not in out
    assert key not in err
    data = _last_line(out)
    assert data["http_error_class"] == "4xx"
    assert data["matched"] == 0
    assert state["counts"].get("e401") == len(gold_doc["tabs"])
    # the status code and the failed-tab count are numbers, so a rate limit is
    # distinguishable from a bad request without reading a body
    assert data["http_status_counts"] == {"401": len(gold_doc["tabs"])}
    assert data["incomplete_tabs"] == len(gold_doc["tabs"])


def test_gold_mode_min_interval_paces_the_calls(gold_dir, capsys):
    """--min-interval spaces call starts apart; the default never sleeps."""
    gold_path, gold_doc = gold_dir
    bodies = _gold_payloads(gold_path, gold_doc)
    interval = 0.3
    with _scripted_server(bodies) as (_server, base, state):
        code = _run(_gold_args(base, gold_path, "--min-interval", str(interval)))
        out, _err = capsys.readouterr()

    assert code == 0
    data = _last_line(out)
    assert data["calls"] >= 2
    assert data["seconds"] >= (data["calls"] - 1) * interval
    assert data["http_status_counts"] == {}
    assert data["incomplete_tabs"] == 0


def test_gold_mode_manifest_mismatch_exits_two(gold_dir, tmp_path, capsys):
    import shutil

    gold_path, gold_doc = gold_dir
    copy = tmp_path / "gold"
    shutil.copytree(gold_path, copy)
    manifest_path = next(copy.glob("gold_manifest_*.json"))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    first_tab = gold_doc["tabs"][0]
    manifest["tab_sha256"][first_tab] = "0" * 64
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    bodies = _gold_payloads(gold_path, gold_doc)
    with _scripted_server(bodies) as (_server, base, _state):
        code = _run(_gold_args(base, copy))
        out, err = capsys.readouterr()

    assert code == 2
    assert out == ""
    assert "sha256" in err
    assert first_tab not in err


def test_gold_mode_writes_record_only_when_asked(gold_dir, tmp_path, capsys):
    gold_path, gold_doc = gold_dir
    bodies = _gold_payloads(gold_path, gold_doc)
    with _scripted_server(bodies) as (_server, base, _state):
        assert _run(_gold_args(base, gold_path)) == 0
        capsys.readouterr()
    assert list(tmp_path.glob("*.json")) == []

    out_path = tmp_path / "record.json"
    with _scripted_server(bodies) as (_server, base, _state):
        assert _run(_gold_args(base, gold_path, "--record-out", str(out_path))) == 0
        capsys.readouterr()
    payload = json.loads(out_path.read_text(encoding="utf-8"))
    assert len(payload["fields"]) == len(gold_doc["fields"])
    assert {f["tab"] for f in payload["fields"]} == set(gold_doc["tabs"])


# --------------------------------------------------------------------------
# --prompt-variant: the L3a lines-contract prompt variants
# --------------------------------------------------------------------------

def _lines_client(gold_path, gold_doc):
    """A provider client that answers each gold tab with its own perfect lines.

    No network: the L3a variants are a prompt change, so the hermetic path can
    only prove the name travels and the plumbing holds, never which variant is
    better (that is the reviewer's live run).
    """
    import gold_lines

    from formextract.resolve import LLMResponse

    texts = gold_lines.GoldLines(gold_path, gold_doc).perfect()

    class Stub(probe._ProviderClient):
        """The probe's own client with its transport replaced by canned lines."""

        def complete(self, prompt, *, model, params):
            tab = prompt.split("Rows of `", 1)[1].split("`", 1)[0]
            text = texts[tab]
            self.calls += 1
            self.finish_counts["stop" if text.endswith("\nend") else "length"] += 1
            return LLMResponse(
                text=text, model=model, params=params, tokens=1, latency_ms=0,
                finish_reason="stop" if text.endswith("\nend") else "length",
            )

    return Stub


def test_gold_mode_lines_contract_prints_the_prompt_variant(gold_dir, capsys, monkeypatch):
    gold_path, gold_doc = gold_dir
    monkeypatch.setattr(probe, "_ProviderClient", _lines_client(gold_path, gold_doc))
    for variant in probe.PROMPT_VARIANTS:
        code = _run(
            _gold_args(
                "http://example.invalid", gold_path,
                "--contract", "lines", "--prompt-variant", variant,
            )
        )
        out, err = capsys.readouterr()
        assert code == 0, (variant, err)
        assert err == ""
        data = _last_line(out)
        assert data["prompt_variant"] == variant
        assert data["precision"] == 1.0
        assert data["recall"] == 1.0
        assert data["gold_fields"] == len(gold_doc["fields"])
        for value in data.values():
            if not isinstance(value, dict):
                _assert_scalar(value)
    # the default is `base`, named in the output like any other
    code = _run(_gold_args("http://example.invalid", gold_path, "--contract", "lines"))
    out, _err = capsys.readouterr()
    assert code == 0
    assert _last_line(out)["prompt_variant"] == "base"


def test_gold_mode_rejects_an_unknown_prompt_variant(gold_dir, capsys):
    gold_path, _gold_doc = gold_dir
    code = _run(
        _gold_args("http://example.invalid", gold_path, "--prompt-variant", "fix3")
    )
    out, err = capsys.readouterr()
    assert code == 2
    assert out == ""
    assert "--prompt-variant" in err


def test_single_tab_mode_prints_the_prompt_variant(tmp_path, capsys):
    """The single-tab path measures the json prompt; it still names the variant."""
    path = _make_workbook(tmp_path / "wb.xlsx")
    with _fake_server(
        ok_body=_openai_body(OK_FIELDS, "stop"), trunc_body=b"", key=""
    ) as (_server, base, _state):
        code = _run(
            [
                str(path),
                "--provider", "openai",
                "--model", "test-model",
                "--base-url", base + "/ok",
                "--env-var", "",
                "--prompt-variant", "fix2",
            ]
        )
        out, err = capsys.readouterr()

    assert code == 0
    assert err == ""
    data = _last_line(out)
    assert data["prompt_variant"] == "fix2"
    assert data["fields"] == 1
    for value in data.values():
        _assert_scalar(value)
