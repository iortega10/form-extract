from __future__ import annotations

import json
import re
from pathlib import Path

from formextract.evals.canned import spec_fragment_response
from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import (
    BBox,
    BindingDraft,
    ControlType,
    ElementRef,
    RegionType,
    ReviewReason,
)
from formextract.pipeline import Pipeline, PipelineConfig
from formextract.resolve import LLMResponse, drafts_to_fields, parse_drafts
from formextract.store import Store


def _ctx(path: Path):
    ing = ingest(path)
    layout = analyze(ing.elements)
    tabs = ing.sheet_names or [f"page {i}" for i in range(ing.page_count or 0)]
    return layout, {e.element_id: e for e in ing.elements}, tabs


def _fields(spec_xlsx):
    layout, elements_by_id, tabs = _ctx(spec_xlsx)
    drafts = parse_drafts(spec_fragment_response())
    return drafts_to_fields(drafts, layout=layout, elements_by_id=elements_by_id, tabs=tabs)


def test_field_id_stable_and_populated(spec_xlsx):
    first = _fields(spec_xlsx)
    second = _fields(spec_xlsx)

    assert [f.field_id for f in first] == [f.field_id for f in second]
    assert all(f.field_id for f in first)
    assert len({f.field_id for f in first}) == len(first)
    assert all(f.source_elements for f in first)
    assert all(f.section_path == ["vendor compliance checklist"] for f in first)
    assert all(f.tab == "Checklist" for f in first)


def test_section_path_from_header_region(spec_xlsx):
    fields = _fields(spec_xlsx)
    assert [f.section_path for f in fields] == [
        ["vendor compliance checklist"],
        ["vendor compliance checklist"],
        ["vendor compliance checklist"],
        ["vendor compliance checklist"],
    ]


def test_field_id_survives_tab_rename(spec_xlsx):
    layout, elements_by_id, _ = _ctx(spec_xlsx)
    drafts = parse_drafts(spec_fragment_response())
    before = drafts_to_fields(drafts, layout=layout, elements_by_id=elements_by_id, tabs=["OldName"])
    after = drafts_to_fields(drafts, layout=layout, elements_by_id=elements_by_id, tabs=["NewName"])

    assert [f.field_id for f in before] == [f.field_id for f in after]
    assert [f.tab for f in before] == ["OldName"] * len(before)
    assert [f.tab for f in after] == ["NewName"] * len(after)


def test_label_edit_retires_only_that_id(spec_xlsx):
    before = _fields(spec_xlsx)
    before_by_label = {f.label_text: f.field_id for f in before}

    layout, elements_by_id, tabs = _ctx(spec_xlsx)
    edited = parse_drafts(spec_fragment_response())
    for draft in edited:
        if draft.label == "Vendor Status":
            draft.label = "Vendor Status (Y/N)"
    after = drafts_to_fields(edited, layout=layout, elements_by_id=elements_by_id, tabs=tabs)
    after_by_label = {f.label_text: f.field_id for f in after}

    assert after_by_label["Vendor Status (Y/N)"] != before_by_label["Vendor Status"]
    for label in ("Which US states do they ship to", "Do they ship to Canada?", "Order Log"):
        assert after_by_label[label] == before_by_label[label]


def test_region_ids_are_content_derived(spec_pdf, spec_xlsx):
    pdf_layout = _ctx(spec_pdf)[0]
    xlsx_layout = _ctx(spec_xlsx)[0]

    pdf_ids = {r.region_id for r in pdf_layout.regions}
    xlsx_ids = {r.region_id for r in xlsx_layout.regions}
    assert pdf_ids == xlsx_ids
    assert all(re.match(r"p\d+:c\d+:(grid|field-row|header|prose):.+", rid) for rid in pdf_ids)


def test_repeated_label_same_tab_distinct_ids(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "repeated.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "Order Log"
    ws["F1"] = "X One"
    ws["A2"] = "Order Log"
    ws["F2"] = "X Two"
    wb.save(path)

    layout, elements_by_id, tabs = _ctx(path)
    region = next(r for r in layout.regions if r.column == 0 and r.type is RegionType.FIELD_ROW)
    drafts = [
        BindingDraft(
            draft_id="d1",
            label="Order Log",
            control_type=ControlType.BOOL,
            bbox=BBox(),
            region_id=region.region_id,
            source_refs=[ElementRef(region.region_id, 0, 0)],
        ),
        BindingDraft(
            draft_id="d2",
            label="Order Log",
            control_type=ControlType.BOOL,
            bbox=BBox(),
            region_id=region.region_id,
            source_refs=[ElementRef(region.region_id, 1, 0)],
        ),
    ]
    fields = drafts_to_fields(drafts, layout=layout, elements_by_id=elements_by_id, tabs=tabs)
    assert len({f.field_id for f in fields}) == 2


def test_unresolvable_element_id_recorded_unknown(spec_xlsx):
    layout, elements_by_id, tabs = _ctx(spec_xlsx)
    label_region = "p0:c0:field-row:which_us_states_do_they_ship_to"
    draft = BindingDraft(
        draft_id="d1",
        label="L",
        control_type=ControlType.TEXT,
        bbox=BBox(),
        region_id=label_region,
        source_refs=[
            ElementRef(label_region, 1, 0),
            ElementRef("no_such_region", 0, 0),
            ElementRef(label_region, 99, 0),
        ],
    )
    (field,) = drafts_to_fields([draft], layout=layout, elements_by_id=elements_by_id, tabs=tabs)

    assert field.source_elements == ["Checklist!3:1"]
    assert field.unresolved_source_refs == [
        "no_such_region:0:0",
        f"{label_region}:99:0",
    ]
    assert field.provenance.review_flag is True
    assert field.provenance.review_reason is ReviewReason.UNRESOLVED_REFERENCE


class _RepeatedTabClient:
    def complete(self, prompt, *, model, params):
        tab = re.search(r"Layout projection for `([^`]+)`", prompt).group(1)
        region = re.search(r"### (p\d+:c0:field-row:[^\s]+)", prompt).group(1)
        band = int(re.search(r"- band (\d+):", prompt).group(1))
        return LLMResponse(
            text=json.dumps(
                {
                    "fields": [
                        {
                            "label": "Order Log",
                            "control_type": "bool",
                            "options": [{"text": "Support Order Log", "selected": True}],
                            "answer": ["true"],
                            "annotations": [],
                            "region_id": region,
                            "source_elements": [
                                {"region_id": region, "band_id": band, "segment_index": 0}
                            ],
                        }
                    ]
                }
            ),
            model=model,
            params=params,
        )


def test_repeated_tabs_distinct_fields(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "six.xlsx"
    wb = Workbook()
    for i in range(6):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = f"Tab{i}"
        ws["A1"] = "Order Log"
        ws["F1"] = "X Support Order Log"
    wb.save(path)

    record = Pipeline(Store(tmp_path / "store"), _RepeatedTabClient(), PipelineConfig()).run(path)

    assert len(record.fields) == 6
    assert {f.tab for f in record.fields} == {f"Tab{i}" for i in range(6)}
    assert len({f.field_id for f in record.fields}) == 6
