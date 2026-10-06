from __future__ import annotations

import json
from pathlib import Path

import pytest

from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import ElementRef
from formextract.pipeline import (
    Pipeline,
    PipelineConfig,
    _resolve_non_answer_element_ids,
    compute_cache_key,
)
from formextract.resolve import LLMResponse
from formextract.store import Store


class _Client:
    def __init__(self, fields):
        self.fields = fields
        self.prompts: list[str] = []

    def complete(self, prompt, *, model, params):
        self.prompts.append(prompt)
        return LLMResponse(
            text=json.dumps({"fields": self.fields}), model=model, params=params, tokens=10
        )


def _refs(layout, elements, texts):
    by_id = {e.element_id: e for e in elements}
    refs = []
    for text in texts:
        target = next(e for e in elements if e.text == text)
        found = None
        for key, bands in layout.bands.items():
            col = key % 1000
            for band_id, band in enumerate(bands):
                for seg, eid in enumerate(band):
                    if eid == target.element_id:
                        region = next(
                            r
                            for r in layout.regions
                            if r.column == col and band_id in r.band_ids
                        )
                        found = ElementRef(
                            region_id=region.region_id,
                            band_id=band_id,
                            segment_index=seg,
                        )
        if found is None:
            raise AssertionError(text)
        refs.append(found)
    return refs


def _field_dict(label, control_type, options, source_texts, annotations=None, answers=None):
    return {
        "label": label,
        "control_type": control_type,
        "options": options,
        "answer": answers or [],
        "annotations": annotations or [],
        "confidence": 0.9,
    }


def _make_workbook(path, rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    for r, (label, cells) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=label)
        for col, text in cells:
            ws[f"{col}{r}"] = text
    wb.save(path)
    wb.close()


def test_x_in_non_answer_column_ignored(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, [("Q", [("F", "Option A"), ("Z", "X")])])
    ing = ingest(path)
    layout = analyze(ing.elements)
    refs = _refs(layout, ing.elements, ["Q", "Option A", "X"])
    field = _field_dict(
        "Q",
        "single_select",
        [{"text": "Option A", "selected": None}],
        [r for r in refs],
        annotations=["X"],
    )
    client = _Client([field])
    pipeline = Pipeline(
        Store(tmp_path / "store"),
        client,
        PipelineConfig(non_answer_columns=["Z"]),
    )
    record = pipeline.run(path)

    assert len(record.fields) == 1
    resolved = record.fields[0]
    assert resolved.value_normalized is None
    assert resolved.provenance.review_flag is False
    assert resolved.ambiguity is None
    assert [o.selected for o in resolved.options] == [None]
    assert "X" in resolved.annotations
    assert "[annotation]" in client.prompts[0]


def test_non_answer_text_is_annotation_not_value(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, [("Q", [("F", "Option A"), ("Z", "Reference only")])])
    ing = ingest(path)
    layout = analyze(ing.elements)
    refs = _refs(layout, ing.elements, ["Q", "Option A", "Reference only"])
    field = _field_dict(
        "Q",
        "single_select",
        [{"text": "Option A", "selected": None}],
        [r for r in refs],
        annotations=["Reference only"],
    )
    client = _Client([field])
    pipeline = Pipeline(
        Store(tmp_path / "store"),
        client,
        PipelineConfig(non_answer_columns=["Z"]),
    )
    record = pipeline.run(path)

    resolved = record.fields[0]
    assert resolved.value_normalized is None
    assert "reference only" in " ".join(a.lower() for a in resolved.annotations)


def test_cache_key_varies_with_non_answer_columns():
    base = {
        "content_hash": "c",
        "pipeline_version": "1",
        "schema_version": "2",
        "prompt_version": "2",
        "model": "m",
        "params": {"temperature": 0},
    }
    a = compute_cache_key(**base, non_answer_columns=("Z",))
    b = compute_cache_key(**base, non_answer_columns=("AA",))
    c = compute_cache_key(**base, non_answer_columns=())
    assert len({a, b, c}) == 3
    # Order and tab-glob shape canonicalise the same way.
    d = compute_cache_key(**base, non_answer_columns=({"tab": "*", "column": "Z"},))
    assert d == a


def test_malformed_non_answer_column_selector_raises(tmp_path):
    store = Store(tmp_path / "store")
    with pytest.raises(ValueError):
        Pipeline(store, None, PipelineConfig(non_answer_columns=["bad!"]))


def test_resolve_non_answer_element_ids_matches_column(tmp_path):
    path = tmp_path / "f.xlsx"
    _make_workbook(path, [("Q", [("F", "Option A"), ("Z", "X")])])
    elements = ingest(path).elements
    ids = _resolve_non_answer_element_ids(elements, ["Z"])
    assert ids == {"Checklist!1:26"}
