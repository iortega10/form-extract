from __future__ import annotations

from pathlib import Path

import pytest

from formextract.evals.synthetic import build
from formextract.store import Store


@pytest.fixture
def spec_pdf(tmp_path: Path) -> Path:
    return build("spec_fragment_pdf", tmp_path / "fixtures")


@pytest.fixture
def spec_xlsx(tmp_path: Path) -> Path:
    return build("spec_fragment_xlsx", tmp_path / "fixtures")


@pytest.fixture
def store(tmp_path: Path) -> Store:
    return Store(tmp_path / "store")


@pytest.fixture
def widths_xlsx(tmp_path: Path) -> Path:
    """Non-uniform column widths: A=1 labels, B=10 empty spacer, C undeclared
    (sheet defaultColWidth=4.0), D:E one range-defined 1-wide entry keyed by
    "D" only, E holds notes; C3:D3 merged spans width units."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Widths"
    ws.sheet_format.defaultColWidth = 4.0
    ws.column_dimensions["A"].width = 1.0
    ws.column_dimensions["B"].width = 10.0
    dim_de = ws.column_dimensions["D"]
    dim_de.width = 1.0
    dim_de.min = 4
    dim_de.max = 5
    ws["A1"] = "Label One"
    ws["C1"] = "Value One"
    ws["E1"] = "Note One"
    ws["A2"] = "Label Two"
    ws["C2"] = "Value Two"
    ws["E2"] = "Note Two"
    ws["C3"] = "Merged"
    ws.merge_cells("C3:D3")
    path = tmp_path / "widths.xlsx"
    wb.save(path)
    return path
