"""The 0.6.0-L1b line contract, end to end against the gold set.

Offline acceptance for the lines contract (the debate doc's L1b row): the
generator's *perfect* canned lines reach strict 1.0 on the gold set through the
real ``Pipeline`` with ``output_contract="lines"``, and every mutation of section
1 moves exactly the counter it should. ``tests/gold_lines.py`` is the bridge that
turns the generator's cell-id-keyed lines into real ``row.seg`` coordinates.

The selection counters carry a documented, pinned gap: the layout classifies a
mark that has several option-like cells to its right as ``competing`` (or
``between`` when one sits to its left), and ``_geometric_mark_decision`` keeps
those short-circuits ahead of the union, so a row that holds several options in
one band (a Yes/No row, a dense multi, a 50-option grid) cannot have its marks
derived from geometry alone. The JSON contract hides this because a model states
``selected`` itself; the lines contract has no model-side ``selected``, so its
``selected_ok`` is lower. ``test_marker_rows_that_geometry_cannot_derive`` pins
the exact numbers so a change is deliberate, and the doc note records it.
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

REPO = Path(__file__).resolve().parents[1]
for _root in (REPO, REPO / "tools"):
    if str(_root) not in sys.path:
        sys.path.insert(0, str(_root))

import gold_lines  # noqa: E402
import make_gold  # noqa: E402
import score_gold  # noqa: E402

from formextract.ingest import ingest  # noqa: E402
from formextract.layout import analyze, strip_marks  # noqa: E402
from formextract.lines import parse_lines  # noqa: E402
from formextract.model import (  # noqa: E402
    BBox,
    BindingDraft,
    ControlType,
    ElementRef,
    InstanceStatus,
    ReviewReason,
)
from formextract.pipeline import (  # noqa: E402
    OUTPUT_CONTRACTS,
    Pipeline,
    PipelineConfig,
    compute_cache_key,
    validate_output_contract,
)
from formextract.resolve import (  # noqa: E402
    LLMResponse,
    ProjectionChunk,
    build_prompt,
    drafts_resolve_cleanly,
    drafts_to_fields,
    parse_lines_response,
    project_chunks,
)
from formextract.reuse import apply_binding  # noqa: E402
from formextract.rows import build_all_rows  # noqa: E402
from formextract.store import Store  # noqa: E402


# --------------------------------------------------------------------------
# Fixtures: build each gold set once per module per gold version, and score the
# perfect lines once. Every version is exercised: the label wording changes in
# v3, and the lines contract must reach strict 1.0 on both.
# --------------------------------------------------------------------------

@pytest.fixture(scope="module", params=make_gold.GOLD_VERSIONS, ids=lambda v: f"v{v}")
def dev(request, tmp_path_factory):
    version = request.param
    base = tmp_path_factory.mktemp(f"lines-dev-v{version}")
    gold, _, _ = make_gold.build_set(
        "dev", make_gold.DEFAULT_SEEDS["dev"], base, gold_version=version
    )
    return gold_lines.GoldLines(base, gold)


@pytest.fixture(scope="module")
def dev_scored(dev, tmp_path_factory):
    fields = dev.run(dev.perfect(), tmp_path_factory.mktemp("dev-perfect"))
    return dev, gold_lines.score(dev.gold, fields)


@pytest.fixture(scope="module", params=make_gold.GOLD_VERSIONS, ids=lambda v: f"v{v}")
def heldout_scored(request, tmp_path_factory):
    version = request.param
    base = tmp_path_factory.mktemp(f"lines-heldout-v{version}")
    gold, _, _ = make_gold.build_set(
        "heldout", make_gold.DEFAULT_SEEDS["heldout"], base, gold_version=version
    )
    gl = gold_lines.GoldLines(base, gold)
    fields = gl.run(gl.perfect(), tmp_path_factory.mktemp("held-perfect"))
    return gl, gold_lines.score(gold, fields)


# --------------------------------------------------------------------------
# Offline acceptance: the perfect lines reach strict 1.0.
# --------------------------------------------------------------------------

def _assert_strict_one(res: dict) -> None:
    overall = res["overall"]
    assert overall["precision"] == 1.0
    assert overall["recall"] == 1.0
    assert overall["matched"] == overall["gold_fields"] == overall["predicted_fields"]
    assert overall["merge_count"] == 0
    assert overall["split_count"] == 0
    assert overall["missed"] == 0
    assert overall["spurious"] == 0
    assert overall["stray_ref_count"] == 0
    assert overall["unresolved_ref_count"] == 0
    assert overall["label_ok"] == overall["label_total"] > 0
    assert overall["options_ok"] == overall["options_total"] > 0
    assert overall["answer_ok"] == overall["answer_total"] > 0
    assert overall["addressed_num"] == overall["addressed_den"] > 0


def test_perfect_lines_reach_strict_one_on_dev(dev_scored):
    dev, res = dev_scored
    assert res["overall"]["gold_fields"] == 244
    _assert_strict_one(res)
    assert set(make_gold.FIELD_TAGS) <= set(res["per_tag"])


def test_perfect_lines_reach_strict_one_on_heldout(heldout_scored):
    _, res = heldout_scored
    _assert_strict_one(res)


def test_canned_lines_artifact_round_trips(dev, tmp_path):
    """The generator's own artifact, written and read back, still scores 1.0."""
    out = tmp_path / "canned"
    make_gold.write_canned_lines(dev.gold, out)
    doc = json.loads((out / "lines_dev.json").read_text(encoding="utf-8"))
    assert doc["set"] == "dev"
    assert set(doc["perfect"]) == set(dev.gold["tabs"])
    assert set(doc["mutations"]) == {
        "drop_line", "wrong_kind", "merge_two", "split_one", "cut_before_end",
    }
    texts = {
        tab: gold_lines.render(doc["perfect"][tab], dev.geometry[tab].refs)
        for tab in dev.gold["tabs"]
    }
    fields = dev.run(texts, tmp_path / "store")
    _assert_strict_one(gold_lines.score(dev.gold, fields))


# --------------------------------------------------------------------------
# Selection: what the injected markers derive, and what they cannot.
# --------------------------------------------------------------------------

def test_stacked_and_typed_selections_are_derived(dev_scored):
    """The T8 union works where one option sits on the marker's own row."""
    _, res = dev_scored
    for tag in ("stacked_multi", "matrix", "typed_value"):
        sub = res["per_tag"][tag]
        assert sub["selected_total"] > 0, tag
        assert sub["selected_ok"] == sub["selected_total"], tag
        assert sub["unexpected_ambiguous"] == 0, tag


def test_marker_rows_that_geometry_cannot_derive(dev_scored):
    """Pin the documented gap: an in-band multi-option row keeps no selection.

    The layout gives such a mark several right candidates (``competing``) or
    candidates on both sides (``between``), and the short-circuits stay ahead of
    the right-only union, so ``_geometric_mark_decision`` cannot pick the option
    the mark belongs to. With no model-side ``selected`` on the lines path the
    field ends up ambiguous and loses every mark. Every number below is the
    measured consequence; the doc note records the cause and the json contrast.
    """
    _, res = dev_scored
    pinned = {
        "dense_multi": (3, 10),
        "grid_50": (0, 10),
        "yes_no_row": (37, 70),
        "gutter_column": (3, 8),
        "non_answer_column": (5, 11),
        "mixed_tab": (54, 90),
    }
    for tag, (ok, total) in pinned.items():
        sub = res["per_tag"][tag]
        assert (sub["selected_ok"], sub["selected_total"]) == (ok, total), tag
    assert res["overall"]["unexpected_ambiguous"] == 42
    assert res["overall"]["selected_ok"] == 100
    assert res["overall"]["selected_total"] == 180


def test_side_by_side_shared_marker_is_the_expected_ambiguity(dev_scored):
    """The gold's ``expected_ambiguous`` pair is not read as a missing convention."""
    _, res = dev_scored
    sub = res["per_tag"]["side_by_side"]
    assert sub["unexpected_ambiguous"] == 0
    assert sub["selected_ambiguous_expected"] == 2
    assert sub["selected_ok"] == 12
    assert sub["selected_total"] == 22
    assert res["overall"]["selected_ambiguous_expected"] == 10


# --------------------------------------------------------------------------
# Mutations: each line-level defect moves exactly the counters it should.
# --------------------------------------------------------------------------

#: Measured against the perfect dev run (``gold_fields`` 244). Only the counters
#: that move are listed; regenerate by rerunning this file.
LINE_MUTATION_DELTAS = {
    "drop_line": {
        "predicted_fields": -1, "matched": -1, "matched_ignoring_kind": -1,
        "missed": 1,
        "label_ok": -1, "label_total": -1, "options_ok": -1, "options_total": -1,
        "selected_ok": -1, "selected_total": -1,
        "selected_ok_under_convention": -1, "addressed_num": -3,
        "precision_num": -1, "precision_den": -1, "recall_num": -1,
    },
    "wrong_kind": {
        "matched": -1, "kind_confusions": 1, "missed": 1, "spurious": 1,
        "label_ok": -1, "label_total": -1, "options_ok": -1, "options_total": -1,
        "selected_ok": -1, "selected_total": -1,
        "selected_ok_under_convention": -1,
        "precision_num": -1, "recall_num": -1,
    },
    "merge_two": {
        "predicted_fields": -1, "matched": -2, "matched_ignoring_kind": -2,
        "missed": 1, "merge_count": 1,
        "label_ok": -2, "label_total": -2, "options_ok": -2, "options_total": -2,
        "selected_ok": -2, "selected_total": -2,
        "selected_ok_under_convention": -2, "addressed_num": -1,
        "precision_num": -2, "precision_den": -2, "recall_num": -2, "recall_den": -1,
    },
    "split_one": {
        "predicted_fields": 1, "matched": -1, "matched_ignoring_kind": -1,
        "split_count": 1,
        "label_ok": -1, "label_total": -1, "options_ok": -1, "options_total": -1,
        "selected_ok": -1, "selected_total": -1,
        "selected_ok_under_convention": -1,
        "precision_num": -1, "precision_den": -1, "recall_num": -1, "recall_den": -1,
    },
    "cut_before_end": {
        # One field per tab is cut: the tail after the last newline is a field
        # line, so the section-3.5 rule drops it.
        "predicted_fields": -16, "matched": -16, "matched_ignoring_kind": -16,
        "missed": 16,
        "label_ok": -16, "label_total": -16, "options_ok": -16, "options_total": -16,
        "selected_ok": -8, "selected_total": -12,
        "selected_ok_under_convention": -7, "selected_ambiguous_expected": -1,
        "answer_ok": -4, "answer_total": -4, "addressed_num": -140,
        "precision_num": -16, "precision_den": -16, "recall_num": -16,
    },
}


def test_line_mutations_move_exactly_the_expected_counters(dev_scored, tmp_path_factory):
    dev, perfect = dev_scored
    baseline = perfect["overall"]
    assert set(LINE_MUTATION_DELTAS) == set(make_gold.line_mutations(dev.gold))
    for name, expected in LINE_MUTATION_DELTAS.items():
        fields = dev.run(dev.mutation(name), tmp_path_factory.mktemp(f"mut-{name}"))
        overall = gold_lines.score(dev.gold, fields)["overall"]
        delta = {
            key: overall[key] - baseline[key]
            for key in score_gold.COUNTER_KEYS
            if isinstance(baseline.get(key), int) and overall.get(key) != baseline[key]
        }
        assert delta == expected, (name, delta)


# --------------------------------------------------------------------------
# Derivation: dispositions, refs, kinds.
# --------------------------------------------------------------------------

def _parse(geo, tab, text, **kwargs):
    chunk = ProjectionChunk(key=tab, text="", page=0)
    response = LLMResponse(
        text=text, model="canned", params={}, finish_reason=kwargs.pop("finish_reason", "stop")
    )
    return parse_lines_response(
        chunk, geo.lattice, response, layout=geo.layout, elements_by_id=geo.by_id, **kwargs
    )


def _resolve(geo, drafts, tab):
    return drafts_to_fields(
        drafts, layout=geo.layout, elements_by_id=geo.by_id, tabs=[tab]
    )


def test_disposition_lines_produce_no_fields(dev):
    tab = dev.gold["tabs"][0]
    geo = dev.geometry[tab]
    text = gold_lines.render(make_gold.perfect_lines(dev.gold)[tab], geo.refs)
    drafts, dispositions, errors, stats = _parse(geo, tab, text)
    assert errors == []
    assert stats.records == len(drafts) + len(dispositions)
    expected_fields = sum(1 for f in dev.gold["fields"] if f["tab"] == tab)
    expected_dispositions = sorted(
        d["disposition"] for d in dev.gold["dispositions"] if d["tab"] == tab
    )
    assert len(drafts) == expected_fields
    assert sorted(d.kind.value for d in dispositions) == expected_dispositions
    assert all(d.rows for d in dispositions)
    assert _resolve(geo, drafts, tab)


def test_a_bad_ref_becomes_an_unresolved_reference(dev):
    tab = dev.gold["tabs"][0]
    geo = dev.geometry[tab]
    drafts, _, errors, stats = _parse(geo, tab, "single L=999.999 O=998.998\nend")
    assert errors == []
    assert stats.records == 1
    fields = _resolve(geo, drafts, tab)
    assert len(fields) == 1
    assert len(fields[0].unresolved_source_refs) == 2
    assert fields[0].source_elements == []
    assert fields[0].provenance.review_reason is ReviewReason.UNRESOLVED_REFERENCE
    assert not drafts_resolve_cleanly(drafts, geo.layout)


def test_an_unknown_kind_is_review_flagged(dev):
    tab = dev.gold["tabs"][0]
    geo = dev.geometry[tab]
    cell = dev.gold["fields"][0]["label_cells"][0]
    drafts, _, errors, stats = _parse(
        geo, tab, f"frobnicate L={geo.refs[cell]}\nend"
    )
    assert errors == []
    assert stats.unknown_kind == 1
    fields = _resolve(geo, drafts, tab)
    assert fields[0].provenance.review_flag is True
    assert fields[0].provenance.review_reason is ReviewReason.AMBIGUOUS_ROLE


def test_a_cut_response_drops_its_tail(dev):
    tab = dev.gold["tabs"][0]
    geo = dev.geometry[tab]
    cell = dev.gold["fields"][0]["label_cells"][0]
    ref = geo.refs[cell]
    good = f"single L={ref}\ntext L={ref}\nend"
    drafts, _, errors, stats = _parse(geo, tab, good)
    assert len(drafts) == 2 and not stats.truncated

    cut = "single L=%s\ntext L=%s" % (ref, ref)
    drafts, dispositions, errors, stats = _parse(geo, tab, cut, finish_reason="length")
    assert stats.truncated and stats.dropped_tail
    assert len(drafts) == 1
    assert errors and errors[0].startswith("response truncated after ")

    empty = _parse(geo, tab, "")
    assert empty[3].empty_response and empty[2] == ["response empty"]


def test_a_cut_response_makes_the_run_partial(dev, tmp_path):
    tab = dev.gold["tabs"][0]
    geo = dev.geometry[tab]
    cell = dev.gold["fields"][0]["label_cells"][0]
    text = "single L=%s\ntext L=%s" % (geo.refs[cell], geo.refs[cell])
    config = PipelineConfig(output_contract="lines")
    record = Pipeline(
        Store(tmp_path / "store"), gold_lines.LinesClient({tab: text}), config
    ).run(dev.gold_dir / f"{tab}.xlsx")
    assert record.status is InstanceStatus.PARTIAL
    assert any("truncated" in error for error in record.errors)
    assert len(record.llm_calls) == 1


# --------------------------------------------------------------------------
# Config, versions, cache key.
# --------------------------------------------------------------------------

def test_output_contract_is_validated(tmp_path):
    assert OUTPUT_CONTRACTS == ("json", "lines")
    validate_output_contract("json")
    validate_output_contract("lines")
    with pytest.raises(ValueError):
        validate_output_contract("yaml")
    with pytest.raises(ValueError):
        Pipeline(Store(tmp_path / "store"), None, PipelineConfig(output_contract="yaml"))


def test_the_cache_key_moves_only_for_a_non_json_contract(tmp_path):
    base = dict(
        content_hash="abc123",
        pipeline_version="4",
        schema_version="2",
        prompt_version="4",
        model="canned",
        params={"temperature": 0},
    )
    default = compute_cache_key(**base)
    assert compute_cache_key(output_contract="json", **base) == default
    lines = compute_cache_key(output_contract="lines", **base)
    assert lines != default
    assert compute_cache_key(output_contract="lines", **base) == lines


def test_the_json_prompt_and_projection_are_unchanged(tmp_path):
    """The 0.5.0 json path is byte-identical: same bytes, same hash."""
    path = tmp_path / "tiny.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Tiny"
    for cell, value in (
        ("A1", "Alpha widget variant?"), ("C1", "X"), ("D1", "Ruby"),
        ("F1", "X"), ("G1", "Teal"),
    ):
        ws[cell] = value
    wb.save(path)
    wb.close()

    ing = ingest(path)
    layout = analyze(ing.elements)
    chunk = project_chunks(layout, ing.elements, ing.sheet_names)[0]
    assert chunk.text == (
        "### p0:c0:field-row:alpha_widget_variant_x_ruby_x_teal [field-row]\n"
        "- band 0: 0=Alpha widget variant? | 1=X | 2=Ruby | 3=X | 4=Teal"
    )
    assert hashlib.sha256(chunk.text.encode()).hexdigest() == (
        "01614ea258964c99df1df330e1c9046f0761d2d3642a9e32c4a4949b1174ee74"
    )
    assert hashlib.sha256(build_prompt(chunk).encode()).hexdigest() == (
        "4397454083fc6032018d81407824dae4007ea5336863971e1f8d96198a5cf099"
    )


# --------------------------------------------------------------------------
# Identity: the label is document text, so a model cannot move a field_id.
# --------------------------------------------------------------------------

def test_field_id_is_stable_across_ref_order(dev):
    target = next(
        f for f in dev.gold["fields"] if len(f["option_cells"]) >= 2
    )
    tab = target["tab"]
    geo = dev.geometry[tab]
    label = geo.refs[target["label_cells"][0]]
    options = [geo.refs[c] for c in target["option_cells"]]
    forward = "multi L=%s O=%s\nend" % (label, ",".join(options))
    backward = "multi L=%s O=%s\nend" % (label, ",".join(reversed(options)))

    first = _resolve(geo, _parse(geo, tab, forward)[0], tab)[0]
    second = _resolve(geo, _parse(geo, tab, backward)[0], tab)[0]
    assert first.label_text == second.label_text
    assert first.field_id == second.field_id
    # The label and the id are the document's; the option list keeps the model's
    # ``O=`` order (the json path's option order is the model's too), so only the
    # option text set is order-independent.
    assert {o.text for o in first.options} == {o.text for o in second.options}


def test_a_two_cell_label_reads_in_document_order(dev):
    target = next(f for f in dev.gold["fields"] if len(f["label_cells"]) == 2)
    tab = target["tab"]
    geo = dev.geometry[tab]
    cells = target["label_cells"]
    expected = " ".join(strip_marks(geo.by_id[c].text) for c in cells)
    assert geo.region_of(cells[0]) == geo.region_of(cells[1])
    refs = [geo.refs[c] for c in cells]

    def label_text(order):
        drafts, _, _, _ = _parse(geo, tab, "text L=%s\nend" % ",".join(order))
        return drafts[0].label

    assert label_text(refs) == expected
    assert label_text(list(reversed(refs))) == expected


# --------------------------------------------------------------------------
# Replay: a text answer is re-read from the target tab's own answer_refs.
# --------------------------------------------------------------------------

def test_a_replayed_text_answer_uses_answer_refs(tmp_path):
    path = tmp_path / "replay.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Tiny"
    ws["A1"] = "Alpha widget variant?"
    ws["B1"] = "Ruby"
    wb.save(path)
    wb.close()

    ing = ingest(path)
    layout = analyze(ing.elements)
    by_id = {e.element_id: e for e in ing.elements}
    lattice = build_all_rows(layout, by_id)[0]
    value = lattice.refs_by_element["Tiny!1:2"]
    value_ref = ElementRef(region_id=value[0], band_id=value[1], segment_index=value[2])

    def draft(**kwargs):
        fields = dict(
            draft_id="d1",
            label="Alpha widget variant?",
            control_type=ControlType.TEXT,
            bbox=BBox(),
            # The label region is the whole row: a one-region layout, where the
            # label-region drop rule cannot tell the label from its value.
            region_id=value_ref.region_id,
            source_refs=[value_ref],
        )
        fields.update(kwargs)
        return BindingDraft(**fields)

    replayed = draft()
    apply_binding([replayed], by_id, layout)
    assert replayed.answers == []

    replayed = draft(answer_refs=(value_ref,))
    apply_binding([replayed], by_id, layout)
    assert replayed.answers == ["Ruby"]


# --------------------------------------------------------------------------
# Totality: no response text can make the lines path raise.
# --------------------------------------------------------------------------

HOSTILE = (
    "",
    "\n\n\n",
    "```json\n{}\n```",
    "single",
    "single L=0.0",
    "single L=0.0 O=0.2\nend",
    "hdr\nnote\nskip\nend",
    'single L="Yes No"\nend',
    'multi L=0.0 O="a"+"b"\nend',
    'text L=\u0660.\u0662\nend',
    "single L=" + "9" * 40 + ".0\nend",
    "single L=0.0 O=0.0-99999999\nend",
    "multi L=0.0 O=12.2-13.1\nend",
    "single L=0.0 O=-1.-1\nend",
    "END.\nsingle L=0.0\nend\nsingle L=0.0",
    "single L=\u0000.0\nend",
    "\U0001f600 L=0.0\nend",
    "a" * 200_000,
    "{\"fields\": []}",
    "single L=0.0 O=0.2" + ",0.2" * 5000 + "\nend",
    "single L:0.0 O:0.2\nend",
    "  SINGLE   l = 0.0 , o = 0.2  \n  END  ",
)


def test_no_response_text_can_make_the_lines_path_raise(dev):
    tab = dev.gold["tabs"][0]
    geo = dev.geometry[tab]
    for text in HOSTILE:
        parsed = parse_lines(text)
        assert isinstance(parsed.errors, list)
        drafts, dispositions, errors, stats = _parse(geo, tab, text)
        assert isinstance(drafts, list)
        assert isinstance(dispositions, list)
        assert isinstance(errors, list)
        assert stats.lines_total >= 0
        _resolve(geo, drafts, tab)


# --------------------------------------------------------------------------
# The reviewer's path: `live_probe.py --gold --contract lines`.
# --------------------------------------------------------------------------

def test_the_probe_gold_runner_can_run_the_lines_contract(dev, monkeypatch):
    import live_probe

    class Stub:
        """A provider client that returns the gold's perfect lines, no network."""

        def complete(self, prompt, *, model, params):
            tab = prompt.split("Rows of `", 1)[1].split("`", 1)[0]
            text = dev.perfect()[tab]
            return LLMResponse(
                text=text, model=model, params=params, tokens=1, latency_ms=0,
                finish_reason="stop" if text.endswith("\nend") else "length",
            )

    monkeypatch.setattr(live_probe, "_ProviderClient", lambda **kwargs: Stub())
    dump, gold_doc, _, _ = live_probe.run_gold_once(
        gold_dir=dev.gold_dir,
        provider="openai",
        base_url="http://example.invalid",
        key="",
        model="canned",
        params={},
        max_tokens=None,
        workers=1,
        contract="lines",
    )
    gold = score_gold.Gold(gold_doc)
    result = score_gold.score_document(
        gold, *score_gold._predicted_fields(gold, dump["fields"]), None
    )
    assert result["overall"]["precision"] == 1.0
    assert result["overall"]["recall"] == 1.0
