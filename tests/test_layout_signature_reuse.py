from __future__ import annotations

import json
import re
from pathlib import Path

from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import (
    ElementRef,
    InstanceStatus,
    ProvenanceSource,
    ReviewReason,
)
from formextract.pipeline import Pipeline, PipelineConfig
from formextract.resolve import LLMResponse
from formextract.store import Store


class _CountingClient:
    def __init__(self, response_text: str):
        self.response_text = response_text
        self.calls = 0

    def complete(self, prompt: str, *, model: str, params: dict) -> LLMResponse:
        self.calls += 1
        return LLMResponse(
            text=self.response_text,
            model=model,
            params=params,
            tokens=1,
            latency_ms=0,
        )


class _PerTabClient:
    def __init__(self, responses: dict[int, str], tabs: list[str]):
        self.responses = responses
        self.tabs = tabs
        self.calls = 0

    def complete(self, prompt: str, *, model: str, params: dict) -> LLMResponse:
        self.calls += 1
        match = re.search(r"Layout projection for `([^`]+)`", prompt)
        tab = match.group(1) if match else (self.tabs[0] if self.tabs else None)
        page = self.tabs.index(tab) if tab in self.tabs else 0
        return LLMResponse(
            text=self.responses[page],
            model=model,
            params=params,
            tokens=1,
            latency_ms=0,
        )


def _make_form(path: Path, tabs: list[str], text_answers: list[str], bool_marked: list[bool]) -> Path:
    """Spec-fragment-shaped tabs plus a per-tab text answer and a per-tab bool mark."""
    from openpyxl import Workbook

    wb = Workbook()
    for i, tab in enumerate(tabs):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tab
        ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
        ws["A3"] = "Which US states do they ship to"
        ws["F3"] = "X"
        ws["G3"] = "All"
        ws["A4"] = "if not all please check those that apply"
        ws["F4"] = "AK AL AR AZ CA CO CT DE FL GA"
        ws["F5"] = "HI IA ID IL IN KS KY LA MA MD"
        ws["A7"] = "Do they ship to Canada?"
        ws["F7"] = "Yes"
        ws["G7"] = "X"
        ws["H7"] = "No"
        ws["A9"] = "Vendor Status"
        ws["F9"] = "Approved"
        ws["G9"] = "X"
        ws["H9"] = "Provisional"
        ws["J9"] = "Confirm on vendor portal"
        ws["A11"] = "Order Log"
        ws.merge_cells("F11:K11")
        ws["F11"] = (
            ("X " if bool_marked[i] else "")
            + "Support Order Log (Check if the vendor is exporting its order log directly)"
        )
        ws["A13"] = "Order Notes"
        ws["F13"] = text_answers[i]
    wb.save(path)
    wb.close()
    return path


def _make_mixed_form(path: Path, tabs: list[str], text_answers: list[str]) -> Path:
    """Alternate two geometries: even tabs plain, odd tabs add a far-right spacer."""
    from openpyxl import Workbook

    wb = Workbook()
    for i, tab in enumerate(tabs):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tab
        ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
        ws["A3"] = "Which US states do they ship to"
        ws["F3"] = "X"
        ws["G3"] = "All"
        ws["A4"] = "if not all please check those that apply"
        ws["F4"] = "AK AL AR AZ CA CO CT DE FL GA"
        ws["F5"] = "HI IA ID IL IN KS KY LA MA MD"
        ws["A7"] = "Do they ship to Canada?"
        ws["F7"] = "Yes"
        ws["G7"] = "X"
        ws["H7"] = "No"
        ws["A9"] = "Vendor Status"
        ws["F9"] = "Approved"
        ws["G9"] = "X"
        ws["H9"] = "Provisional"
        ws["J9"] = "Confirm on vendor portal"
        ws["A11"] = "Order Log"
        ws.merge_cells("F11:K11")
        ws["F11"] = "X Support Order Log (Check if the vendor is exporting its order log directly)"
        ws["A13"] = "Order Notes"
        ws["F13"] = text_answers[i]
        if i % 2 == 1:
            ws["Z20"] = "spacer"
    wb.save(path)
    wb.close()
    return path


def _analyze(path: Path):
    elements = ingest(path).elements
    return analyze(elements), elements


def _element_ref(layout, page: int, element_id: str) -> ElementRef:
    for key, bands in layout.bands.items():
        if key // 1000 != page:
            continue
        column = key % 1000
        for band_id, band in enumerate(bands):
            for segment_index, eid in enumerate(band):
                if eid == element_id:
                    region = next(
                        r
                        for r in layout.regions
                        if r.bbox.page == page
                        and r.column == column
                        and band_id in r.band_ids
                    )
                    return ElementRef(
                        region_id=region.region_id,
                        band_id=band_id,
                        segment_index=segment_index,
                    )
    raise AssertionError(f"no layout ref for {element_id!r} on page {page}")


def _label_element(elements, page: int, text: str):
    for e in elements:
        if e.bbox.page == page and e.text == text:
            return e
    raise AssertionError(f"missing label {text!r} on page {page}")


def _value_element(elements, page: int, label):
    for e in elements:
        if e.bbox.page != page or e.element_id == label.element_id:
            continue
        if e.bbox.y0 == label.bbox.y0 and e.bbox.x0 > label.bbox.x0:
            return e
    raise AssertionError(f"missing value for {label.text!r} on page {page}")


def _ref_dict(ref: ElementRef) -> dict:
    return {
        "region_id": ref.region_id,
        "band_id": ref.band_id,
        "segment_index": ref.segment_index,
    }


def _text_field(layout, elements, page: int) -> dict:
    label = _label_element(elements, page, "Order Notes")
    value = _value_element(elements, page, label)
    lref = _element_ref(layout, page, label.element_id)
    vref = _element_ref(layout, page, value.element_id)
    return {
        "label": "Order Notes",
        "control_type": "text",
        "options": [],
        "answer": [value.text],
        "annotations": [],
        "region_id": lref.region_id,
        "source_elements": [_ref_dict(lref), _ref_dict(vref)],
        "confidence": 0.9,
    }


def _text_field_label_only(layout, elements, page: int) -> dict:
    label = _label_element(elements, page, "Order Notes")
    lref = _element_ref(layout, page, label.element_id)
    return {
        "label": "Order Notes",
        "control_type": "text",
        "options": [],
        "answer": [],
        "annotations": [],
        "region_id": lref.region_id,
        "source_elements": [_ref_dict(lref)],
        "confidence": 0.9,
    }


def _bool_field(layout, elements, page: int) -> dict:
    label = _label_element(elements, page, "Order Log")
    value = _value_element(elements, page, label)
    lref = _element_ref(layout, page, label.element_id)
    vref = _element_ref(layout, page, value.element_id)
    return {
        "label": "Order Log",
        "control_type": "bool",
        "options": [{"text": "Support Order Log", "selected": True}],
        "answer": ["true"],
        "annotations": [],
        "region_id": lref.region_id,
        "source_elements": [_ref_dict(lref), _ref_dict(vref)],
        "confidence": 0.9,
    }


def _form_responses(path: Path, tabs: list[str]) -> dict[int, str]:
    layout, elements = _analyze(path)
    return {
        page: json.dumps(
            {"fields": [_text_field(layout, elements, page), _bool_field(layout, elements, page)]}
        )
        for page in range(len(tabs))
    }


def _text_responses(path: Path, tabs: list[str]) -> dict[int, str]:
    layout, elements = _analyze(path)
    return {
        page: json.dumps({"fields": [_text_field(layout, elements, page)]})
        for page in range(len(tabs))
    }


def test_reuse_never_replays_another_tabs_answers(tmp_path: Path):
    tabs = [f"Checklist {i}" for i in range(6)]
    path = _make_form(
        tmp_path / "form.xlsx",
        tabs,
        text_answers=tabs,
        bool_marked=[i % 2 == 0 for i in range(6)],
    )
    client = _PerTabClient(_form_responses(path, tabs), tabs)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    assert record.status is InstanceStatus.COMPLETE
    assert client.calls == 1
    assert len(record.fields) == 12  # 2 fields x 6 tabs

    text_fields = [f for f in record.fields if f.label_text == "Order Notes"]
    bool_fields = [f for f in record.fields if f.label_text == "Order Log"]

    # Every reused text field carries its own tab's answer, never the exemplar's.
    assert {f.value_normalized for f in text_fields} == set(tabs)
    assert [f.value_normalized for f in text_fields] == [f.tab for f in text_fields]

    # Every reused bool field carries its own tab's mark.
    for f in bool_fields:
        idx = tabs.index(f.tab)
        if idx % 2 == 0:
            assert f.value_normalized == "true"
            assert [o.selected for o in f.options] == [True]
        else:
            assert f.value_normalized is None
            assert [o.selected for o in f.options] == [None]

    sources = [f.provenance.source for f in record.fields]
    assert sources.count(ProvenanceSource.LLM) == 2
    assert sources.count(ProvenanceSource.REPLAY) == 10


def test_reuse_reads_text_answers_per_tab(tmp_path: Path):
    tabs = ["East", "West", "Central"]
    answers = ["alpha", "beta", "gamma"]
    path = _make_form(
        tmp_path / "text.xlsx",
        tabs,
        text_answers=answers,
        bool_marked=[True, True, True],
    )
    client = _PerTabClient(_text_responses(path, tabs), tabs)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    assert record.status is InstanceStatus.COMPLETE
    assert client.calls == 1
    fields = [f for f in record.fields if f.label_text == "Order Notes"]
    assert {f.tab: f.value_normalized for f in fields} == {
        "East": "alpha",
        "West": "beta",
        "Central": "gamma",
    }
    assert all(f.provenance.source is ProvenanceSource.REPLAY for f in fields if f.tab != "East")


def test_empty_answer_cell_yields_none_with_flag(tmp_path: Path):
    tabs = ["Tab A", "Tab B"]
    path = _make_form(
        tmp_path / "empty.xlsx",
        tabs,
        text_answers=["alpha", " "],
        bool_marked=[True, True],
    )
    client = _PerTabClient(_text_responses(path, tabs), tabs)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    assert record.status is InstanceStatus.COMPLETE
    fields = {f.tab: f for f in record.fields if f.label_text == "Order Notes"}
    assert fields["Tab A"].value_normalized == "alpha"
    assert fields["Tab B"].value_normalized is None
    assert fields["Tab B"].provenance.source is ProvenanceSource.REPLAY
    assert fields["Tab B"].provenance.review_flag is True
    assert fields["Tab B"].provenance.review_reason is ReviewReason.REPLAY_MISMATCH


def _make_label_suffixed_form(path: Path, tabs: list[str]) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    for i, tab in enumerate(tabs):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tab
        ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
        ws["A3"] = "Which US states do they ship to"
        ws["F3"] = "X"
        ws["G3"] = "All"
        ws["A4"] = "if not all please check those that apply"
        ws["F4"] = "AK AL AR AZ CA CO CT DE FL GA"
        ws["F5"] = "HI IA ID IL IN KS KY LA MA MD"
        ws["A7"] = "Do they ship to Canada?"
        ws["F7"] = "Yes"
        ws["G7"] = "X"
        ws["H7"] = "No"
        ws["A13"] = f"Order Notes {i}"
        ws["F13"] = f"answer {i}"
    wb.save(path)
    wb.close()
    return path


def test_same_geometry_different_labels_no_reuse(tmp_path: Path):
    tabs = [f"Tab {i}" for i in range(6)]
    path = _make_label_suffixed_form(tmp_path / "labels.xlsx", tabs)
    response = json.dumps(
        {
            "fields": [
                {
                    "label": "Order Notes",
                    "control_type": "text",
                    "options": [],
                    "answer": ["v"],
                    "annotations": [],
                }
            ]
        }
    )
    client = _CountingClient(response)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    assert client.calls == 6
    assert record.fields
    assert all(f.provenance.source is ProvenanceSource.LLM for f in record.fields)


def test_reuse_miss_falls_back_per_tab(tmp_path: Path):
    path = _make_form(
        tmp_path / "miss.xlsx",
        ["Tab A", "Tab B"],
        text_answers=["alpha", "beta"],
        bool_marked=[True, True],
    )
    bad_response = json.dumps(
        {
            "fields": [
                {
                    "label": "Bad",
                    "control_type": "text",
                    "options": [],
                    "answer": [],
                    "annotations": [],
                    "region_id": "p0:c0:field-row:nope",
                    "source_elements": [
                        {
                            "region_id": "p0:c0:field-row:nope",
                            "band_id": 0,
                            "segment_index": 0,
                        }
                    ],
                    "confidence": 0.5,
                }
            ]
        }
    )
    client = _CountingClient(bad_response)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    assert client.calls == 2
    assert record.fields
    assert all(f.provenance.source is ProvenanceSource.LLM for f in record.fields)


def test_binding_address_miss_refuses_whole_tab(tmp_path: Path):
    tabs = [f"Tab {i}" for i in range(6)]
    path = tmp_path / "miss-address.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    for i, tab in enumerate(tabs):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tab
        ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
        ws["A3"] = "Which US states do they ship to"
        ws["F3"] = "X"
        ws["G3"] = "All"
        ws["A4"] = "if not all please check those that apply"
        ws["F4"] = "AK AL AR AZ CA CO CT DE FL GA"
        ws["F5"] = "HI IA ID IL IN KS KY LA MA MD"
        ws["A7"] = "Do they ship to Canada?"
        ws["F7"] = "Yes"
        ws["G7"] = "X"
        ws["H7"] = "No"
        ws["A9"] = "Vendor Status"
        ws["F9"] = "Approved"
        ws["G9"] = "X"
        ws["H9"] = "Provisional"
        ws["J9"] = "Confirm on vendor portal"
        ws["A11"] = "Order Log"
        ws.merge_cells("F11:K11")
        ws["F11"] = "X Support Order Log (Check if the vendor is exporting its order log directly)"
        ws["A13"] = "Order Notes"
        if i != 4:
            ws["F13"] = f"answer {i}"
    wb.save(path)
    wb.close()

    layout, elements = _analyze(path)
    responses = {}
    for page in range(len(tabs)):
        if page == 4:
            responses[page] = json.dumps({"fields": [_text_field_label_only(layout, elements, page)]})
        else:
            responses[page] = json.dumps({"fields": [_text_field(layout, elements, page)]})

    client = _PerTabClient(responses, tabs)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    assert record.status is InstanceStatus.COMPLETE
    assert client.calls == 2  # exemplar tab 0 + the refused tab 4
    by_tab: dict[str, list] = {}
    for f in record.fields:
        by_tab.setdefault(f.tab, []).append(f.provenance.source)
    # Tabs 1, 2, 3, 5 reused the exemplar; tab 4 was refused wholesale.
    for i, tab in enumerate(tabs):
        sources = by_tab.get(tab, [])
        if i in (0, 4):
            assert sources and all(s is ProvenanceSource.LLM for s in sources)
        else:
            assert sources and all(s is ProvenanceSource.REPLAY for s in sources)


def test_reuse_never_crosses_signature_groups(tmp_path: Path):
    tabs = [f"Tab {i}" for i in range(6)]
    path = _make_mixed_form(tmp_path / "mixed.xlsx", tabs, [f"v{i}" for i in range(6)])
    client = _PerTabClient(_text_responses(path, tabs), tabs)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    assert record.status is InstanceStatus.COMPLETE
    assert client.calls == 2
    sources = [f.provenance.source for f in record.fields]
    assert sources.count(ProvenanceSource.LLM) == 2
    assert sources.count(ProvenanceSource.REPLAY) == 4


def test_reused_field_ids_equal_fresh_authoring_ids(tmp_path: Path):
    tabs = ["One", "Two", "Three"]
    path = _make_form(
        tmp_path / "ids.xlsx",
        tabs,
        text_answers=["a", "b", "c"],
        bool_marked=[True, True, True],
    )
    responses = _form_responses(path, tabs)

    on_client = _PerTabClient(responses, tabs)
    on_record = Pipeline(
        Store(tmp_path / "store-on"), on_client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    off_client = _PerTabClient(responses, tabs)
    off_record = Pipeline(
        Store(tmp_path / "store-off"), off_client, PipelineConfig(reuse_layout_bindings=False)
    ).run(path)

    def structural(field):
        return (field.tab, field.field_id, tuple(field.source_elements))

    assert {structural(f) for f in on_record.fields} == {
        structural(f) for f in off_record.fields
    }
    assert on_client.calls == 1
    assert off_client.calls == 3


def test_author_failure_does_not_starve_reuse_tabs(tmp_path: Path):
    tabs = [f"Tab {i}" for i in range(4)]
    path = _make_form(
        tmp_path / "fail.xlsx",
        tabs,
        text_answers=["a", "b", "c", "d"],
        bool_marked=[True, True, True, True],
    )

    class _FlakyClient:
        def __init__(self):
            self.calls = 0
            self.prompts: list[str] = []

        def complete(self, prompt: str, *, model: str, params: dict) -> LLMResponse:
            self.calls += 1
            self.prompts.append(prompt)
            raise RuntimeError("transport down")

    client = _FlakyClient()
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    # The exemplar fails, but every other tab still falls back to its own
    # authoring attempt (4 distinct chunks, 3 transport retries each).
    assert len({p for p in client.prompts}) == 4
    assert client.calls == 12
    assert record.status is InstanceStatus.PARTIAL
    assert any("transport failed" in e for e in record.errors)


def test_reuse_replay_provenance_and_flag(tmp_path: Path):
    tabs = ["A", "B", "C"]
    path = _make_form(
        tmp_path / "prov.xlsx",
        tabs,
        text_answers=["x", "y", "z"],
        bool_marked=[True, True, True],
    )
    client = _PerTabClient(_text_responses(path, tabs), tabs)
    record = Pipeline(
        Store(tmp_path / "store"), client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    replayed = [f for f in record.fields if f.provenance.source is ProvenanceSource.REPLAY]
    assert len(replayed) == 2
    for f in replayed:
        assert f.provenance.review_flag is False
        assert f.provenance.review_reason is None
        assert f.value_normalized is not None


def test_reuse_is_per_run_and_in_memory(tmp_path: Path):
    tabs = ["A", "B"]
    path = _make_form(
        tmp_path / "mem.xlsx",
        tabs,
        text_answers=["a", "b"],
        bool_marked=[True, True],
    )
    responses = _text_responses(path, tabs)

    # A single run reuses within itself: one exemplar call for two tabs, and
    # the exemplar map lives only on the Pipeline instance for that run.
    store1 = Store(tmp_path / "store1")
    client1 = _PerTabClient(responses, tabs)
    rec1 = Pipeline(store1, client1, PipelineConfig(reuse_layout_bindings=True)).run(path)
    assert client1.calls == 1

    # A fresh Pipeline + fresh store re-authors the exemplar: nothing about the
    # exemplar map was persisted that the second run could read back.
    store2 = Store(tmp_path / "store2")
    client2 = _PerTabClient(responses, tabs)
    rec2 = Pipeline(store2, client2, PipelineConfig(reuse_layout_bindings=True)).run(path)
    assert client2.calls == 1

    for store in (store1, store2):
        names = {p.name for p in store.root.iterdir()}
        assert "reuse" not in names
        assert "exemplar" not in names

    assert [f.value_normalized for f in rec1.fields] == [
        f.value_normalized for f in rec2.fields
    ]


def test_two_runs_identical_modulo_run_ids(tmp_path: Path):
    tabs = ["A", "B", "C"]
    path = _make_form(
        tmp_path / "tworuns.xlsx",
        tabs,
        text_answers=["a", "b", "c"],
        bool_marked=[True, False, True],
    )
    responses = _form_responses(path, tabs)

    rec1 = Pipeline(
        Store(tmp_path / "s1"), _PerTabClient(responses, tabs), PipelineConfig(reuse_layout_bindings=True)
    ).run(path)
    rec2 = Pipeline(
        Store(tmp_path / "s2"), _PerTabClient(responses, tabs), PipelineConfig(reuse_layout_bindings=True)
    ).run(path)

    def comparable(rec):
        return [
            (
                f.tab,
                f.label_text,
                f.field_id,
                f.value_normalized,
                tuple(f.source_elements),
                f.provenance.source.value if f.provenance.source else None,
                f.provenance.review_flag,
            )
            for f in rec.fields
        ]

    assert comparable(rec1) == comparable(rec2)


def test_call_counts_off_vs_on(tmp_path: Path):
    tabs = [f"Tab {i}" for i in range(6)]
    path = _make_form(
        tmp_path / "counts.xlsx",
        tabs,
        text_answers=[f"v{i}" for i in range(6)],
        bool_marked=[True] * 6,
    )
    responses = _text_responses(path, tabs)

    on_client = _PerTabClient(responses, tabs)
    Pipeline(
        Store(tmp_path / "on"), on_client, PipelineConfig(reuse_layout_bindings=True)
    ).run(path)
    off_client = _PerTabClient(responses, tabs)
    Pipeline(
        Store(tmp_path / "off"), off_client, PipelineConfig(reuse_layout_bindings=False)
    ).run(path)

    off = off_client.calls
    on = on_client.calls
    assert off == 6
    assert on == 1
    assert on < off

    mixed = _make_mixed_form(tmp_path / "mixed-counts.xlsx", tabs, [f"v{i}" for i in range(6)])
    mixed_responses = _text_responses(mixed, tabs)
    mixed_on = _PerTabClient(mixed_responses, tabs)
    Pipeline(
        Store(tmp_path / "mixed-on"), mixed_on, PipelineConfig(reuse_layout_bindings=True)
    ).run(mixed)
    mixed_off = _PerTabClient(mixed_responses, tabs)
    Pipeline(
        Store(tmp_path / "mixed-off"), mixed_off, PipelineConfig(reuse_layout_bindings=False)
    ).run(mixed)

    mixed_off_count = mixed_off.calls
    mixed_on_count = mixed_on.calls
    assert mixed_off_count == 6
    assert mixed_on_count == 2
    assert mixed_on_count < mixed_off_count
