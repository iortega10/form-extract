from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

from formextract.ingest import ingest
from formextract.layout import _page_geometry_payload, analyze, page_geometry_signature


def _elements(path: Path):
    return ingest(path).elements


def _layout(path: Path):
    return analyze(_elements(path))


def _signature(path: Path, page: int = 0) -> str:
    elements = _elements(path)
    return page_geometry_signature(_layout(path), {e.element_id: e for e in elements}, page)


def _base_workbook(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
    ws["A3"] = "Which US states do they ship to"
    ws["F3"] = "X"
    ws["G3"] = "All"
    ws["A4"] = "if not all please check those that apply"
    ws["F4"] = "AK AL AR AZ CA CO CT DE FL GA"
    ws["F5"] = "HI IA ID IL IN KS KY LA MA MD"
    ws["A7"] = "Do they ship to Canada?"
    ws["F7"] = "Yes"
    ws["G7"] = "X"
    ws["H7"] = "No"
    ws["A9"] = "Vendor Status"
    ws["F9"] = "Approved"
    ws["G9"] = "X"
    ws["H9"] = "Provisional"
    ws["J9"] = "Confirm on vendor portal"
    wb.save(path)
    wb.close()
    return path


def test_signature_ignores_text_and_region_ids(tmp_path: Path):
    p1 = _base_workbook(tmp_path / "a.xlsx")
    p2 = tmp_path / "b.xlsx"

    from openpyxl import load_workbook

    wb = load_workbook(p1)
    ws = wb.active
    ws.title = "Renamed"
    ws["A3"] = "Which US states do they ship to CHANGED"
    ws["G3"] = "All CHANGED"
    wb.save(p2)
    wb.close()

    # Same geometry, different text and different tab name -> same signature.
    assert _signature(p1, 0) == _signature(p2, 0)
    # Region ids are content-derived, so they do differ.
    assert {r.region_id for r in _layout(p1).regions} != {
        r.region_id for r in _layout(p2).regions
    }


def test_signature_changes_with_geometry(tmp_path: Path):
    from openpyxl import load_workbook

    base = _base_workbook(tmp_path / "base.xlsx")
    base_sig = _signature(base, 0)

    # Move a band to a different row.
    moved = tmp_path / "moved.xlsx"
    wb = load_workbook(base)
    ws = wb.active
    ws["A7"] = None
    ws["F7"] = None
    ws["G7"] = None
    ws["H7"] = None
    ws["A8"] = "Do they ship to Canada?"
    ws["F8"] = "Yes"
    ws["G8"] = "X"
    ws["H8"] = "No"
    wb.save(moved)
    wb.close()
    assert _signature(moved, 0) != base_sig

    # Add a merged range.
    merged = tmp_path / "merged.xlsx"
    wb = load_workbook(base)
    ws = wb.active
    ws.merge_cells("A15:C15")
    ws["A15"] = "Merged"
    wb.save(merged)
    wb.close()
    assert _signature(merged, 0) != base_sig

    # Add a row (and therefore a band).
    added = tmp_path / "added.xlsx"
    wb = load_workbook(base)
    ws = wb.active
    ws["A11"] = "Extra Field"
    ws["F11"] = "Value"
    wb.save(added)
    wb.close()
    assert _signature(added, 0) != base_sig

    # Widen a column class.
    widened = tmp_path / "widened.xlsx"
    wb = load_workbook(base)
    ws = wb.active
    ws.column_dimensions["A"].width = 40.0
    wb.save(widened)
    wb.close()
    assert _signature(widened, 0) != base_sig


def test_signature_is_deterministic_across_runs(tmp_path: Path):
    path = _base_workbook(tmp_path / "det.xlsx")
    script = (
        "import sys; from pathlib import Path; "
        "from formextract.ingest import ingest; "
        "from formextract.layout import analyze, page_geometry_signature; "
        "p = Path(sys.argv[1]); el = ingest(p).elements; lo = analyze(el); "
        "print(page_geometry_signature(lo, {e.element_id: e for e in el}, 0))"
    )
    outs = []
    for seed in ("1", "424242"):
        result = subprocess.run(
            [sys.executable, "-c", script, str(path)],
            capture_output=True,
            text=True,
            env={**os.environ, "PYTHONHASHSEED": seed},
        )
        assert result.returncode == 0, result.stderr
        outs.append(result.stdout.strip())
    assert outs[0] == outs[1]
    assert outs[0] == _signature(path, 0)


def test_signature_independent_of_tab_name_and_position(tmp_path: Path):
    from openpyxl import Workbook

    def write_tab(ws, variant: int) -> None:
        ws["A1"] = "VENDOR COMPLIANCE CHECKLIST"
        ws["A3"] = "Which US states do they ship to"
        ws["F3"] = "X"
        ws["G3"] = "All"
        if variant == 1:
            ws["A5"] = "Extra"
            ws["F5"] = "Value"
        elif variant == 2:
            ws["A7"] = "Another"
            ws["F7"] = "Field"

    def make(path: Path, order: list[int]) -> Path:
        wb = Workbook()
        for i, variant in enumerate(order):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = f"Sheet {i}"
            write_tab(ws, variant)
        wb.save(path)
        wb.close()
        return path

    p1 = make(tmp_path / "order1.xlsx", [0, 1, 2])
    p2 = make(tmp_path / "order2.xlsx", [2, 0, 1])

    def signatures(path: Path) -> set[str]:
        elements = _elements(path)
        layout = _layout(path)
        by_id = {e.element_id: e for e in elements}
        return {
            page_geometry_signature(layout, by_id, page)
            for page in range(layout.page_count)
        }

    assert signatures(p1) == signatures(p2)


def test_signature_payload_is_integers_only(tmp_path: Path):
    path = _base_workbook(tmp_path / "payload.xlsx")
    elements = _elements(path)
    layout = _layout(path)
    by_id = {e.element_id: e for e in elements}

    # The payload is quantised at the layout's banding scale (0.6 * median
    # element height, `int(round(value / tol_y))`), so it carries only ints
    # and region-type strings — no floats, no text, no region ids.
    payload = _page_geometry_payload(layout, by_id, 0)
    canonical = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    assert "." not in canonical
    assert page_geometry_signature(layout, by_id, 0) == hashlib.sha256(
        canonical.encode("utf-8")
    ).hexdigest()
