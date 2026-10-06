"""0.6.0-T4: row coverage, the opt-in min_coverage status rule, and the cache key."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import pytest

from formextract import coverage as coverage_module
from formextract import pipeline as pipeline_module
from formextract import resolve as resolve_module
from formextract.coverage import UNBOUND_CAP
from formextract.evals.canned import spec_fragment_response
from formextract.evals.synthetic import build
from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import (
    InstanceRecord,
    InstanceStatus,
    RegionType,
    TabCoverage,
    UnboundUnit,
    from_json,
    to_json,
)
from formextract.pipeline import (
    Pipeline,
    PipelineConfig,
    compute_cache_key,
    validate_min_coverage,
)
from formextract.resolve import LLMResponse
from formextract.store import Store, canonical_json

README = Path(__file__).resolve().parents[1] / "README.md"

# compute_cache_key's output for _CACHE_BASE under the 0.5.0 formula; the test
# recomputes that formula independently, so this pins bytes and not code.
RECORDED_0_5_0_KEY = "88ae54407b54ff054d36f2e0c3c5249d951cfe55c7967b09077640b8e9f5e185"
_CACHE_BASE = dict(
    content_hash="c",
    pipeline_version="3",
    schema_version="2",
    prompt_version="3",
    model="m",
    params={"temperature": 0},
)

# spec_fragment_xlsx: five anchor rows in column 0 plus one GRID region in column 1.
_SPEC_LABEL_REGION = "p0:c0:field-row:which_us_states_do_they_ship_to"
_SPEC_GRID_REGION = "p0:c1:grid:ak_al_ar_az_ca_co_ct_de_fl_ga"


# --------------------------------------------------------------------------- #
# builders
# --------------------------------------------------------------------------- #
def _rows_workbook(path: Path, rows: int = 98, *, value_column: bool = True) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    for i in range(rows):
        ws.cell(row=i + 1, column=1, value=f"Label Row {i:02d}")
        if value_column:
            ws.cell(row=i + 1, column=6, value="X")
    wb.save(path)
    wb.close()
    return path


def _rows_per_tab_workbook(path: Path, rows_by_tab: dict[str, int]) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    for i, (tab, rows) in enumerate(rows_by_tab.items()):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tab
        for r in range(rows):
            ws.cell(row=r + 1, column=1, value=f"{tab} Row {r:02d}")
            ws.cell(row=r + 1, column=6, value="X")
    wb.save(path)
    wb.close()
    return path


def _identical_tabs_workbook(path: Path, names: tuple[str, ...], rows: int) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    for i, tab in enumerate(names):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = tab
        for r in range(rows):
            ws.cell(row=r + 1, column=1, value=f"Common Row {r:02d}")
            ws.cell(row=r + 1, column=6, value="X")
    wb.save(path)
    wb.close()
    return path


def _hidden_first_workbook(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    hidden = wb.active
    hidden.title = "Hidden"
    hidden["A1"] = "Field One"
    hidden["F1"] = "v1"
    hidden.sheet_state = "hidden"
    for name in ("Second", "Third"):
        ws = wb.create_sheet(name)
        ws["A1"] = "Field One"
        ws["F1"] = "v1"
        ws["A2"] = "Field Two"
        ws["F2"] = "v2"
    wb.save(path)
    wb.close()
    return path


# --------------------------------------------------------------------------- #
# clients and small helpers
# --------------------------------------------------------------------------- #
def _raw_field(label: str, region_id: str, band_id: int, segment_index: int = 0) -> dict:
    return {
        "label": label,
        "control_type": "text",
        "options": [],
        "answer": ["v"],
        "annotations": [],
        "region_id": region_id,
        "source_elements": [
            {
                "region_id": region_id,
                "band_id": band_id,
                "segment_index": segment_index,
            }
        ],
    }


class _Client:
    """A canned client: one fixed response body, with a call counter."""

    def __init__(self, body):
        self.body = body if isinstance(body, str) else json.dumps({"fields": body})
        self.calls = 0

    def complete(self, prompt, *, model, params):
        self.calls += 1
        return LLMResponse(text=self.body, model=model, params=params, tokens=7, latency_ms=0)


class _PromptRegionClient:
    """Cites band 0 of the first region the prompt lists, once per tab."""

    def __init__(self, label: str = "Field One"):
        self.label = label
        self.calls = 0

    def complete(self, prompt, *, model, params):
        self.calls += 1
        match = re.search(r"### (\S+) \[", prompt)
        fields = []
        if match:
            fields.append(_raw_field(self.label, match.group(1), 0))
        return LLMResponse(
            text=json.dumps({"fields": fields}),
            model=model,
            params=params,
            tokens=7,
            latency_ms=0,
        )


def _run(path, client, config=None, *, store_name="store"):
    store = Store(Path(path).parent / store_name)
    return Pipeline(store, client, config or PipelineConfig()).run(path)


def _layout_of(path):
    elements = ingest(path).elements
    return analyze(elements), elements


def _label_region_id(layout) -> str:
    return next(
        r.region_id
        for r in layout.regions
        if r.type is RegionType.FIELD_ROW and r.column == 0
    )


def _old_050_key() -> str:
    """The pre-0.6.0 compute_cache_key, reimplemented from its source."""
    params_hash = hashlib.sha256(canonical_json({"temperature": 0}).encode()).hexdigest()
    empty = hashlib.sha256(canonical_json([]).encode()).hexdigest()
    parts = "|".join(["c", "3", "2", "3", "m", "0", params_hash, empty, empty, "0", "0"])
    return hashlib.sha256(parts.encode()).hexdigest()


def _consumed_observation(path, tmp_path, config=None):
    """units_consumed of the one tab from a fresh run (fresh store each time).

    The single field cites band 1 of the label region with an out-of-range
    segment index, so the ref cannot resolve: the correct reading is 0.
    """
    index = len(list(tmp_path.glob("obs-*")))
    store = Store(tmp_path / f"obs-{index}")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    body = [_raw_field("Dangling", region, 1, segment_index=999)]
    record = Pipeline(store, _Client(body), config or PipelineConfig()).run(path)
    assert record.fields[0].source_elements == []
    return record.coverage[0].units_consumed


# --------------------------------------------------------------------------- #
# the block
# --------------------------------------------------------------------------- #
def test_full_coverage_gives_ratio_one(tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=98)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    fields = [_raw_field(f"Row {i:02d}", region, i) for i in range(98)]

    record = _run(path, _Client(fields))

    (block,) = record.coverage
    assert block.tab == "Checklist"
    assert (block.units_total, block.units_consumed, block.ratio) == (98, 98, 1.0)
    assert (block.anchors, block.grid_regions) == (98, 0)
    assert block.chunked is False
    assert block.max_anchors_per_tab == 98
    assert block.unbound == []
    assert block.unbound_truncated is False
    assert record.coverage_min_ratio == 1.0


def test_five_of_ninety_eight_anchor_rows(tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=98)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    fields = [_raw_field(f"Row {i:02d}", region, i) for i in range(5)]

    record = _run(path, _Client(fields))

    (block,) = record.coverage
    assert (block.units_total, block.units_consumed, block.ratio) == (98, 5, 0.051)
    assert block.unbound_truncated is True
    assert len(block.unbound) == UNBOUND_CAP == 50
    # deterministic order (page, column, band); handles only, no text
    assert (block.unbound[0].region_id, block.unbound[0].band_id) == (region, 5)
    assert (block.unbound[-1].region_id, block.unbound[-1].band_id) == (region, 54)
    assert record.status is InstanceStatus.COMPLETE
    assert record.errors == []


def test_grid_unit_consumed_or_unbound(tmp_path: Path):
    path = build("spec_fragment_xlsx", tmp_path / "fixtures")

    unconsumed = _run(path, _Client(spec_fragment_response()), store_name="s-off")
    (block,) = unconsumed.coverage
    assert (block.units_total, block.units_consumed, block.ratio) == (6, 4, 0.6667)
    assert (block.anchors, block.grid_regions) == (5, 1)
    assert UnboundUnit(region_id=_SPEC_GRID_REGION, band_id=1) in block.unbound
    assert UnboundUnit(region_id=_SPEC_LABEL_REGION, band_id=2) in block.unbound

    # Any one cited element anywhere in the region consumes the whole unit: the
    # GRID unit at band 1 stands for both of its bands (1 and 2).
    body = [
        _raw_field("Which states", _SPEC_LABEL_REGION, 1),
        _raw_field("States", _SPEC_GRID_REGION, 2),
    ]
    consumed = _run(path, _Client(body), store_name="s-on")
    (block,) = consumed.coverage
    assert (block.units_total, block.units_consumed, block.ratio) == (6, 2, 0.3333)
    assert all(u.region_id != _SPEC_GRID_REGION for u in block.unbound)


def test_non_answer_columns_exclusion(tmp_path: Path):
    path = build("spec_fragment_xlsx", tmp_path / "fixtures")

    # Column F is a reference column: its GRID unit leaves the denominator
    # while the column-A anchor rows stay.
    with_ref = _run(
        path,
        _Client(spec_fragment_response()),
        PipelineConfig(non_answer_columns=["F"]),
        store_name="s-ref",
    )
    (block,) = with_ref.coverage
    assert (block.units_total, block.anchors, block.grid_regions) == (5, 5, 0)
    assert block.ratio == 0.8

    # A tab whose only column is non-answer has zero units: ineligible.
    only_labels = _rows_workbook(tmp_path / "a.xlsx", rows=10, value_column=False)
    layout, _ = _layout_of(only_labels)
    region = _label_region_id(layout)
    empty = _run(
        only_labels,
        _Client([_raw_field("Row 00", region, 0)]),
        PipelineConfig(non_answer_columns=["A"], min_coverage=0.9),
        store_name="s-empty",
    )
    (block,) = empty.coverage
    assert (block.units_total, block.anchors, block.grid_regions) == (0, 0, 0)
    assert block.ratio is None
    assert empty.coverage_min_ratio is None
    assert empty.errors == []
    assert empty.status is InstanceStatus.COMPLETE
    assert with_ref.coverage_min_ratio == 0.8


def test_unresolved_refs_do_not_count(tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=10)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    record = _run(path, _Client([_raw_field("Gone", region, 999)]))

    assert record.fields[0].unresolved_source_refs == [f"{region}:999:0"]
    assert record.fields[0].source_elements == []
    (block,) = record.coverage
    assert block.units_consumed == 0
    assert block.ratio == 0.0
    assert len(block.unbound) == 10


def test_record_without_llm_has_empty_coverage(tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=4)
    record = Pipeline(Store(tmp_path / "store"), None, PipelineConfig()).run(path)
    assert record.coverage == []
    assert record.coverage_min_ratio is None


# --------------------------------------------------------------------------- #
# tabs
# --------------------------------------------------------------------------- #
def test_hidden_first_sheet_uses_true_tab_names(tmp_path: Path):
    path = _hidden_first_workbook(tmp_path / "wb.xlsx")
    record = _run(path, _PromptRegionClient())

    assert record.tabs == ["Second", "Third"]
    assert record.hidden_sheets == ["Hidden"]
    assert [block.tab for block in record.coverage] == ["Second", "Third"]


def test_per_tab_and_record_level_minimum(tmp_path: Path):
    path = _rows_per_tab_workbook(tmp_path / "wb.xlsx", {"East": 2, "West": 4})
    record = _run(path, _PromptRegionClient())

    by_tab = {block.tab: block for block in record.coverage}
    assert (by_tab["East"].units_total, by_tab["East"].ratio) == (2, 0.5)
    assert (by_tab["West"].units_total, by_tab["West"].ratio) == (4, 0.25)
    assert record.coverage_min_ratio == 0.25


def test_chunk_workers_agree(tmp_path: Path):
    path = _rows_per_tab_workbook(tmp_path / "wb.xlsx", {"East": 3, "West": 5})
    serial = _run(path, _PromptRegionClient(), PipelineConfig(chunk_workers=1), store_name="s1")
    pooled = _run(path, _PromptRegionClient(), PipelineConfig(chunk_workers=4), store_name="s2")

    assert [to_json(b) for b in serial.coverage] == [to_json(b) for b in pooled.coverage]
    assert serial.coverage_min_ratio == pooled.coverage_min_ratio


def test_reuse_path_counts_replayed_tabs(tmp_path: Path):
    path = _identical_tabs_workbook(tmp_path / "wb.xlsx", ("East", "West"), rows=3)
    client = _PromptRegionClient()
    record = _run(path, client, PipelineConfig(reuse_layout_bindings=True))

    assert client.calls == 1  # West is replayed, not authored
    assert [block.tab for block in record.coverage] == ["East", "West"]
    east, west = record.coverage
    assert (east.units_total, east.units_consumed, east.ratio) == (3, 1, 0.3333)
    assert (west.units_total, west.units_consumed, west.ratio) == (3, 1, 0.3333)
    assert record.coverage_min_ratio == 0.3333


# --------------------------------------------------------------------------- #
# the min_coverage status rule
# --------------------------------------------------------------------------- #
def test_min_coverage_error_string_and_partial_status(tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=98)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    fields = [_raw_field(f"Row {i:02d}", region, i) for i in range(5)]

    record = _run(path, _Client(fields), PipelineConfig(min_coverage=0.5))

    assert record.status is InstanceStatus.PARTIAL
    assert record.errors == ["tab Checklist: row coverage 0.05 < 0.50"]


def test_status_unchanged_without_min_coverage(tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=98)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    fields = [_raw_field(f"Row {i:02d}", region, i) for i in range(5)]

    record = _run(path, _Client(fields))

    assert record.status is InstanceStatus.COMPLETE
    assert record.errors == []
    assert record.coverage[0].ratio == 0.051


def test_only_tabs_below_the_threshold_get_an_error(tmp_path: Path):
    path = _rows_per_tab_workbook(tmp_path / "wb.xlsx", {"East": 2, "West": 4})
    record = _run(
        path,
        _PromptRegionClient(),
        PipelineConfig(min_coverage=0.4),
        store_name="s-threshold",
    )

    assert record.status is InstanceStatus.PARTIAL
    assert record.errors == ["tab West: row coverage 0.25 < 0.40"]


def test_min_coverage_rerun_costs_no_tokens(tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=98)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    fields = [_raw_field(f"Row {i:02d}", region, i) for i in range(5)]
    store = Store(tmp_path / "store")

    first = _Client(fields)
    record = Pipeline(store, first, PipelineConfig(min_coverage=0.5)).run(path)
    assert first.calls == 1
    assert record.status is InstanceStatus.PARTIAL

    # A PARTIAL record is not served from the instance cache, but the chunk's
    # response parsed without field errors, so it is served from the call
    # cache: the rerun recomputes and spends nothing.
    second = _Client(fields)
    again = Pipeline(store, second, PipelineConfig(min_coverage=0.5)).run(path)
    assert second.calls == 0
    assert again.status is InstanceStatus.PARTIAL
    assert [to_json(b) for b in again.coverage] == [to_json(b) for b in record.coverage]


def test_field_error_response_is_never_cached_so_a_rerun_re_spends(tmp_path: Path):
    # A response the parser accepted but that carried field-level errors is
    # archived unindexed (resolve.py), so an ordinary rerun re-sends it rather
    # than serving it from the call cache.
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=4)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    bad = {**_raw_field("Broken", region, 0), "control_type": "nonsense"}
    store = Store(tmp_path / "store")

    first = _Client([bad])
    record = Pipeline(store, first, PipelineConfig()).run(path)
    assert first.calls == 1
    assert record.status is InstanceStatus.PARTIAL
    assert any("field 0" in err for err in record.errors), record.errors

    second = _Client([bad])
    again = Pipeline(store, second, PipelineConfig()).run(path)
    assert second.calls == 1  # re-sent, not served from the call cache
    assert again.status is InstanceStatus.PARTIAL


def test_min_coverage_validator(tmp_path: Path):
    validate_min_coverage(None)
    validate_min_coverage(1)
    validate_min_coverage(0.05)
    for bad in (True, False, "0.5", 0, 0.0, -0.1, 1.01, float("nan")):
        with pytest.raises(ValueError) as excinfo:
            validate_min_coverage(bad)
        assert "min_coverage" in str(excinfo.value)

    with pytest.raises(ValueError) as excinfo:
        Pipeline(Store(tmp_path / "store"), None, PipelineConfig(min_coverage=True))
    assert "min_coverage" in str(excinfo.value)


# --------------------------------------------------------------------------- #
# cache key
# --------------------------------------------------------------------------- #
def test_cache_key_unchanged_by_default():
    assert _old_050_key() == RECORDED_0_5_0_KEY
    assert compute_cache_key(**_CACHE_BASE) == RECORDED_0_5_0_KEY
    assert compute_cache_key(min_coverage=None, **_CACHE_BASE) == RECORDED_0_5_0_KEY


def test_cache_key_min_coverage_knob():
    default = compute_cache_key(**_CACHE_BASE)
    half = compute_cache_key(min_coverage=0.5, **_CACHE_BASE)
    half_again = compute_cache_key(min_coverage=0.50, **_CACHE_BASE)
    quarter = compute_cache_key(min_coverage=0.25, **_CACHE_BASE)

    assert half != default
    assert half == half_again
    assert quarter != half
    assert quarter != default


# --------------------------------------------------------------------------- #
# serialization and determinism
# --------------------------------------------------------------------------- #
def test_record_round_trips_and_old_record_loads(tmp_path: Path):
    path = build("spec_fragment_xlsx", tmp_path / "fixtures")
    record = _run(path, _Client(spec_fragment_response()))

    rehydrated = from_json(InstanceRecord, to_json(record))
    assert rehydrated.coverage == record.coverage
    assert rehydrated.coverage_min_ratio == record.coverage_min_ratio
    assert isinstance(rehydrated.coverage[0], TabCoverage)
    assert isinstance(rehydrated.coverage[0].unbound[0], UnboundUnit)

    # A record written before the block existed (no keys at all) still loads.
    legacy = json.loads(to_json(record))
    assert "coverage" in legacy
    del legacy["coverage"]
    del legacy["coverage_min_ratio"]
    old = from_json(InstanceRecord, json.dumps(legacy))
    assert old.coverage == []
    assert old.coverage_min_ratio is None


def test_two_runs_are_byte_identical(tmp_path: Path):
    path = build("spec_fragment_xlsx", tmp_path / "fixtures")
    first = _run(path, _Client(spec_fragment_response()), store_name="s1")
    second = _run(path, _Client(spec_fragment_response()), store_name="s2")

    assert to_json(first.coverage) == to_json(second.coverage)
    assert first.coverage_min_ratio == second.coverage_min_ratio


# --------------------------------------------------------------------------- #
# layout invariant
# --------------------------------------------------------------------------- #
def test_grid_regions_never_contain_header_bands(tmp_path: Path):
    paths = [
        build("spec_fragment_xlsx", tmp_path / "f1"),
        build("spec_fragment_pdf", tmp_path / "f2"),
    ]
    for path in paths:
        layout, _ = _layout_of(path)
        assert any(r.type is RegionType.GRID for r in layout.regions), path
        by_column: dict[tuple[int, int], list[tuple[list[int], RegionType]]] = {}
        for region in layout.regions:
            by_column.setdefault((region.bbox.page, region.column), []).append(
                (region.band_ids, region.type)
            )
        for groups in by_column.values():
            header_bands = {
                band
                for band_ids, rtype in groups
                if rtype is RegionType.HEADER
                for band in band_ids
            }
            for band_ids, rtype in groups:
                # every region is one run of consecutive same-type bands
                assert band_ids == list(range(band_ids[0], band_ids[-1] + 1))
                if rtype is RegionType.GRID:
                    assert not (header_bands & set(band_ids))


# --------------------------------------------------------------------------- #
# README
# --------------------------------------------------------------------------- #
def test_readme_documents_row_coverage():
    text = README.read_text(encoding="utf-8")

    assert "**Row coverage.**" in text
    assert "not a\nquality score" in text
    assert "no default" in text
    assert "coverage_min_ratio" in text
    assert "unbound_truncated" in text
    assert "an example threshold, not a recommendation" in text


# --------------------------------------------------------------------------- #
# mutation cases, each with the anti-vacuity triple
# --------------------------------------------------------------------------- #
def _patch(monkeypatch, target, name, replacement):
    """Replace `target.name`, assert the seam exists, and log every call."""
    assert hasattr(target, name), f"seam {name!r} is missing"
    calls: list = []

    def patched(*args, **kwargs):
        calls.append(True)
        return replacement(*args, **kwargs)

    monkeypatch.setattr(target, name, patched)
    assert getattr(target, name) is patched
    return calls


def test_mutation_counting_unresolved_refs_as_consumed(monkeypatch, tmp_path: Path):
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=4)
    correct = _consumed_observation(path, tmp_path)
    assert correct == 0

    def resolve_dangling_refs(layout, fields):
        ids = {eid for f in fields or [] for eid in f.source_elements}
        for field in fields or []:
            for ref in field.unresolved_source_refs:
                region_id, band_id, _seg = ref.rsplit(":", 2)
                region = next(r for r in layout.regions if r.region_id == region_id)
                bands = layout.bands.get(region.bbox.page * 1000 + region.column, [])
                ids.update(bands[int(band_id)])
        return ids

    calls = _patch(
        monkeypatch, coverage_module, "_consumed_element_ids", resolve_dangling_refs
    )
    observed = _consumed_observation(path, tmp_path)

    assert calls, "the mutation was never reached"
    assert observed == 1
    assert observed != correct


def test_mutation_excluding_grid_units(monkeypatch, tmp_path: Path):
    path = build("spec_fragment_xlsx", tmp_path / "fixtures")
    correct = _run(path, _Client(spec_fragment_response()), store_name="s-correct")
    assert correct.coverage[0].grid_regions == 1

    real_units = coverage_module._units_for_page

    def anchors_only(layout, page):
        return [unit for unit in real_units(layout, page) if not unit[3]]

    calls = _patch(monkeypatch, coverage_module, "_units_for_page", anchors_only)
    observed = _run(path, _Client(spec_fragment_response()), store_name="s-mutated")

    assert calls, "the mutation was never reached"
    assert observed.coverage[0].grid_regions == 0
    assert observed.coverage[0].units_total == 5
    assert observed.coverage[0].grid_regions != correct.coverage[0].grid_regions


def test_mutation_counting_non_answer_units(monkeypatch, tmp_path: Path):
    path = build("spec_fragment_xlsx", tmp_path / "fixtures")
    config = PipelineConfig(non_answer_columns=["F"])
    correct = _run(path, _Client(spec_fragment_response()), config, store_name="s-correct")
    assert correct.coverage[0].units_total == 5

    calls = _patch(monkeypatch, coverage_module, "_is_excluded", lambda elements, na: False)
    observed = _run(path, _Client(spec_fragment_response()), config, store_name="s-mutated")

    assert calls, "the mutation was never reached"
    assert observed.coverage[0].units_total == 6
    assert observed.coverage[0].units_total != correct.coverage[0].units_total


def test_mutation_ratio_none_treated_as_one(monkeypatch, tmp_path: Path):
    path = _rows_workbook(tmp_path / "a.xlsx", rows=10, value_column=False)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    config = PipelineConfig(non_answer_columns=["A"])

    correct = _run(
        path, _Client([_raw_field("Row 00", region, 0)]), config, store_name="s-correct"
    )
    assert correct.coverage[0].ratio is None
    assert correct.coverage_min_ratio is None

    calls = _patch(
        monkeypatch,
        coverage_module,
        "_ratio",
        lambda consumed, total: (
            round(consumed / total, 4) if total else 1.0
        ),
    )
    observed = _run(
        path, _Client([_raw_field("Row 00", region, 0)]), config, store_name="s-mutated"
    )

    assert calls, "the mutation was never reached"
    assert observed.coverage[0].ratio == 1.0
    assert observed.coverage_min_ratio == 1.0
    assert observed.coverage_min_ratio != correct.coverage_min_ratio


def test_mutation_min_coverage_dropped_from_the_key(monkeypatch, tmp_path: Path):
    # With min_coverage in the key, a COMPLETE run at threshold 1.0 (all rows
    # consumed) and a later run at the default threshold are different keys, so
    # the second run re-authors. Dropping min_coverage from the key collapses
    # them: the second run is served the first run's instance and never calls
    # the model (a silently stale COMPLETE).
    path = _rows_workbook(tmp_path / "wb.xlsx", rows=4)
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)
    fields = [_raw_field(f"Row {i:02d}", region, i) for i in range(4)]

    def key_of(config, store_name):
        store = Store(tmp_path / store_name)
        client = _Client(fields)
        record = Pipeline(store, client, config).run(path)
        return store, client, record

    _, _, first = key_of(PipelineConfig(min_coverage=1.0), "s-correct")
    assert first.status is InstanceStatus.COMPLETE
    _, _, second = key_of(PipelineConfig(), "s-correct")
    # different key -> not served from the instance cache -> re-authored
    assert second.instance_id != first.instance_id

    real = pipeline_module.compute_cache_key

    def ignoring_min_coverage(**kwargs):
        kwargs.pop("min_coverage", None)
        return real(**kwargs)

    calls = _patch(
        monkeypatch, pipeline_module, "compute_cache_key", ignoring_min_coverage
    )
    _, _, mutated_first = key_of(PipelineConfig(min_coverage=1.0), "s-mutated")
    _, mutated_client, mutated_second = key_of(PipelineConfig(), "s-mutated")

    assert calls, "the mutation was never reached"
    assert mutated_client.calls == 0  # served a stale COMPLETE, no model call
    assert mutated_second.instance_id == mutated_first.instance_id
