"""0.6.0 T1 (retry hygiene) and T1b (force re-spends) tests."""
from __future__ import annotations

import enum
import hashlib
import json
import threading
from pathlib import Path

import pytest

from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import InstanceStatus
from formextract.pipeline import Pipeline, PipelineConfig
from formextract.resolve import (
    LLMResponse,
    NonRetryable,
    ProjectionChunk,
    _complete_with_retries,
    author_drafts,
    build_prompt,
)

PARAMS = {"temperature": 0}

# T1 and T1b promise no prompt change. These are the sha256 (hex) of the UTF-8
# prompt bytes `build_prompt(ProjectionChunk(key="tab", text="content"))`
# produces, keyed by `include_address`. Produced on the pre-T1/T2 commit
# dc166a1 in a `git worktree` of that commit with the venv python, then
# re-checked on this tree; the prompt-producing code (`BINDING_OUTPUT_CONTRACT`,
# `_binding_contract`, `build_prompt`) is byte-identical between the two.
PINNED_PROMPT_SHA256 = {
    False: "3e9ec5acb821c45cc2732cb3325176be7e282cfb8a17b62501855c61d69e8c8c",
    True: "0a6f55dbf1a79477fcf9e6907f9b9fde4d6d135d8b25ee3c0632103b9f383b73",
}


def test_prompt_bytes_are_unchanged():
    chunk = ProjectionChunk(key="tab", text="content")
    for include_address, expected in PINNED_PROMPT_SHA256.items():
        text = build_prompt(chunk, include_address=include_address)
        assert hashlib.sha256(text.encode("utf-8")).hexdigest() == expected, (
            f"prompt drift for include_address={include_address}"
        )


class _StatusError(Exception):
    def __init__(self, status_code: int) -> None:
        super().__init__(f"HTTP {status_code}")
        self.status_code = status_code


class _CodeError(Exception):
    def __init__(self, code) -> None:
        super().__init__("coded")
        self.code = code


class _Code(enum.IntEnum):
    BAD_REQUEST = 400


class _Response:
    status_code = 400


class _WrappedError(Exception):
    def __init__(self) -> None:
        super().__init__("wrapped")
        self.response = _Response()


class _HostileError(Exception):
    def __getattr__(self, name):  # pragma: no cover - exercised via the loop
        raise RuntimeError(f"attribute {name} is booby-trapped")


class _RaisingClient:
    def __init__(self, make_exc) -> None:
        self.make_exc = make_exc
        self.calls = 0

    def complete(self, prompt, *, model, params):
        self.calls += 1
        raise self.make_exc()


def _run(client, *, attempts: int = 3):
    return _complete_with_retries(
        client, "prompt", model="m", params=PARAMS, attempts=attempts, backoff_seconds=0
    )


@pytest.mark.parametrize(
    "status,expected",
    [(400, 1), (401, 1), (403, 1), (404, 1), (408, 3), (429, 3), (500, 3), (503, 3)],
)
def test_4xx_costs_one_call_and_408_429_retry(status, expected):
    client = _RaisingClient(lambda: _StatusError(status))
    with pytest.raises(_StatusError):
        _run(client)
    assert client.calls == expected


def test_plain_exception_still_retries_to_the_limit():
    client = _RaisingClient(lambda: RuntimeError("boom"))
    with pytest.raises(RuntimeError):
        _run(client)
    assert client.calls == 3


def test_non_retryable_marker_costs_one_call_and_keeps_the_reason():
    client = _RaisingClient(lambda: NonRetryable("the provider said no"))
    with pytest.raises(NonRetryable) as excinfo:
        _run(client)
    assert client.calls == 1
    assert "the provider said no" in str(excinfo.value)


def test_code_attribute_digit_string_is_non_retryable():
    client = _RaisingClient(lambda: _CodeError("400"))
    with pytest.raises(_CodeError):
        _run(client)
    assert client.calls == 1


def test_code_attribute_int_enum_is_non_retryable():
    client = _RaisingClient(lambda: _CodeError(_Code.BAD_REQUEST))
    with pytest.raises(_CodeError):
        _run(client)
    assert client.calls == 1


def test_response_status_code_is_non_retryable():
    client = _RaisingClient(_WrappedError)
    with pytest.raises(_WrappedError):
        _run(client)
    assert client.calls == 1


def test_hostile_getattr_does_not_crash_the_loop():
    client = _RaisingClient(_HostileError)
    with pytest.raises(_HostileError):
        _run(client)
    assert client.calls == 3


def _two_tab_workbook(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    good = wb.active
    good.title = "Good"
    good["A1"] = "Field Good"
    bad = wb.create_sheet("Bad")
    bad["A1"] = "Field Bad"
    wb.save(path)
    wb.close()
    return path


def _field(label: str) -> dict:
    return {
        "label": label,
        "control_type": "text",
        "options": [],
        "answer": ["v"],
        "annotations": [],
    }


def test_pipeline_4xx_costs_one_call_for_that_chunk(tmp_path: Path):
    path = _two_tab_workbook(tmp_path / "two.xlsx")
    store_root = tmp_path / "store"

    class Client:
        def __init__(self) -> None:
            self.prompts: list[str] = []

        def complete(self, prompt, *, model, params):
            self.prompts.append(prompt)
            if "Layout projection for `Bad`" in prompt:
                raise _StatusError(400)
            return LLMResponse(
                text=json.dumps({"fields": [_field("Good")]}), model=model, params=params
            )

    from formextract.store import Store

    client = Client()
    record = Pipeline(Store(store_root), client, PipelineConfig()).run(path)

    assert record.status is InstanceStatus.PARTIAL
    assert [f.label_text for f in record.fields] == ["Good"]
    assert sum("Layout projection for `Bad`" in p for p in client.prompts) == 1
    assert any(
        "tab Bad: transport failed: HTTP 400" in err for err in record.errors
    ), record.errors


# --- T1b: force=True re-spends ---------------------------------------------


class _OnceClient:
    def __init__(self, text: str) -> None:
        self.text = text
        self.calls = 0
        self.prompts: list[str] = []

    def complete(self, prompt, *, model, params):
        self.calls += 1
        self.prompts.append(prompt)
        return LLMResponse(text=self.text, model=model, params=params)


def test_force_resends_unchanged_prompt_and_replaces_the_index(store):
    chunk = ProjectionChunk(key="tab", text="content")
    first = _OnceClient(json.dumps({"fields": [_field("A")]}))
    author_drafts([chunk], first, model="m", params=PARAMS, store=store)
    assert first.calls == 1

    second = _OnceClient(json.dumps({"fields": [_field("A")]}))
    author_drafts([chunk], second, model="m", params=PARAMS, store=store)
    assert second.calls == 0  # unchanged prompt served from the call cache

    forced = _OnceClient(json.dumps({"fields": [_field("B")]}))
    drafts, _, _ = author_drafts(
        [chunk], forced, model="m", params=PARAMS, store=store, force=True
    )
    assert forced.calls == 1  # force skips the call-cache lookup
    assert [d.label for d in drafts] == ["B"]

    after = _OnceClient(json.dumps({"fields": [_field("B")]}))
    drafts, _, _ = author_drafts(
        [chunk], after, model="m", params=PARAMS, store=store
    )
    assert after.calls == 0  # the fresh response replaced the indexed one
    assert [d.label for d in drafts] == ["B"]


def test_force_with_chunk_workers_makes_one_call_per_chunk(store):
    chunks = [ProjectionChunk(key=f"tab{i}", text=f"chunk {i}") for i in range(4)]

    class Client:
        def __init__(self) -> None:
            self.calls = 0
            self.lock = threading.Lock()

        def complete(self, prompt, *, model, params):
            with self.lock:
                self.calls += 1
            return LLMResponse(
                text=json.dumps({"fields": [_field("F")]}), model=model, params=params
            )

    forced = Client()
    author_drafts(
        chunks, forced, model="m", params=PARAMS, store=store, max_workers=4, force=True
    )
    assert forced.calls == 4

    cached = Client()
    author_drafts(chunks, cached, model="m", params=PARAMS, store=store, max_workers=4)
    assert cached.calls == 0


def test_force_respects_the_call_budget(store):
    chunks = [
        ProjectionChunk(key="a", text="ca"),
        ProjectionChunk(key="b", text="cb"),
    ]
    client = _OnceClient(json.dumps({"fields": [_field("F")]}))
    drafts, calls, errors = author_drafts(
        chunks,
        client,
        model="m",
        params=PARAMS,
        store=store,
        force=True,
        call_budget=1,
        backoff_seconds=0,
    )
    assert client.calls == 1
    assert any("call budget exhausted" in err.reason for err in errors), errors


def _identical_two_tab(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    for i, name in enumerate(("A", "B")):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws["A1"] = "Reviewed by"
        ws["B1"] = "J. Rivera"
    wb.save(path)
    wb.close()
    return path


def _label_ref(path: Path, page: int, text: str) -> tuple[str, dict]:
    elements = ingest(path).elements
    layout = analyze(elements)
    target = next(e for e in elements if e.bbox.page == page and e.text == text)
    for key, bands in layout.bands.items():
        if key // 1000 != page:
            continue
        column = key % 1000
        for band_id, band in enumerate(bands):
            for segment_index, eid in enumerate(band):
                if eid != target.element_id:
                    continue
                region = next(
                    r
                    for r in layout.regions
                    if r.bbox.page == page
                    and r.column == column
                    and band_id in r.band_ids
                )
                return region.region_id, {
                    "region_id": region.region_id,
                    "band_id": band_id,
                    "segment_index": segment_index,
                }
    raise AssertionError(f"no layout ref for {text!r} on page {page}")


def test_reuse_path_honours_force(tmp_path: Path):
    path = _identical_two_tab(tmp_path / "same.xlsx")
    from formextract.store import Store

    store = Store(tmp_path / "store")

    def response_for(page: int) -> str:
        region_id, ref = _label_ref(path, page, "Reviewed by")
        return json.dumps(
            {
                "fields": [
                    {
                        "label": "Reviewed by",
                        "control_type": "text",
                        "options": [],
                        "answer": ["J. Rivera"],
                        "annotations": [],
                        "region_id": region_id,
                        "source_elements": [ref],
                    }
                ]
            }
        )

    class Client:
        def __init__(self) -> None:
            self.calls = 0

        def complete(self, prompt, *, model, params):
            self.calls += 1
            page = 0 if "`A`" in prompt else 1
            return LLMResponse(text=response_for(page), model=model, params=params)

    config = PipelineConfig(reuse_layout_bindings=True)

    first = Client()
    record = Pipeline(store, first, config).run(path)
    assert first.calls == 1  # one exemplar; the identical tab replayed
    assert record.status is InstanceStatus.COMPLETE

    cached = Client()
    Pipeline(store, cached, config).run(path)
    assert cached.calls == 0

    forced = Client()
    Pipeline(store, forced, config).run(path, force=True)
    assert forced.calls == 1  # the exemplar is re-authored, not served from cache
