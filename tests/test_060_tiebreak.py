"""0.6.0-T4: the deterministic tie-break for drafts that share an anchor band.

Before 0.6.0 two drafts citing the same anchor band kept the model's response
order (``list.sort`` is stable), so ``ordinal`` — and therefore ``field_id`` —
flipped when the model emitted them the other way round. The secondary key in
``_anchor_sort_key`` closes that hole. These tests are the before/after
``field_id`` evidence the plan calls for.
"""
from __future__ import annotations

import json
from pathlib import Path

from formextract.ingest import ingest
from formextract.layout import analyze
from formextract.model import BBox, BindingDraft, ControlType, ElementRef, RegionType
from formextract.resolve import (
    _anchor_sort_key,
    drafts_to_fields,
    parse_drafts_with_errors,
)
from formextract.resolve import LLMResponse
from formextract.pipeline import Pipeline, PipelineConfig
from formextract.store import Store

REGION = "p0:c0:field-row:tie_workbook"


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _workbook(path: Path, rows: int = 3) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    for i in range(rows):
        ws.cell(row=i + 1, column=1, value=f"Row {i:02d}")
        ws.cell(row=i + 1, column=6, value="X")
    wb.save(path)
    wb.close()
    return path


def _layout_of(path: Path):
    elements = ingest(path).elements
    layout = analyze(elements)
    return layout, elements


def _label_region_id(layout) -> str:
    return next(
        r.region_id
        for r in layout.regions
        if r.type is RegionType.FIELD_ROW and r.column == 0
    )


def _draft(label: str, region_id: str, band: int, segment: int = 0, answer: str = "v"):
    return BindingDraft(
        draft_id=f"d-{label}-{band}-{segment}",
        label=label,
        control_type=ControlType.TEXT,
        bbox=BBox(),
        answers=[answer],
        region_id=region_id,
        source_refs=[
            ElementRef(region_id=region_id, band_id=band, segment_index=segment)
        ],
    )


def _labels_and_ids(path: Path, drafts: list[BindingDraft]):
    """(labels, field_ids, fields) for `drafts` resolved against `path`'s layout."""
    layout, elements = _layout_of(path)
    fields = drafts_to_fields(
        drafts,
        layout=layout,
        elements_by_id={e.element_id: e for e in elements},
        tabs=["Checklist"],
    )
    return [f.label_text for f in fields], [f.field_id for f in fields], fields


def _old_key_order(drafts: list[BindingDraft], region_id: str):
    """Pre-0.6.0: sort by anchor band only, stable on the emission order."""
    from formextract.resolve import _anchor_band_id

    return [
        draft.label
        for draft in sorted(
            drafts,
            key=lambda d: (
                _anchor_band_id(d, region_id)
                if _anchor_band_id(d, region_id) is not None
                else 10**9
            ),
        )
    ]


# --------------------------------------------------------------------------- #
# unit-level: the secondary key
# --------------------------------------------------------------------------- #
def test_tied_anchor_band_is_broken_by_label(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    beta = _draft("Beta", region, 0)
    alpha = _draft("Alpha", region, 0)

    labels_a, ids_a, _ = _labels_and_ids(path, [beta, alpha])
    labels_b, ids_b, _ = _labels_and_ids(path, [alpha, beta])

    assert labels_a == ["Alpha", "Beta"]
    assert labels_b == ["Alpha", "Beta"]
    assert ids_a == ids_b


def test_secondary_key_prefers_smallest_band_segment(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    # Cited (band, segment) beats the label: "Zzz"@(0,0) sorts before "Aaa"@(0,1).
    high = _draft("Aaa", region, 0, segment=1)
    low = _draft("Zzz", region, 0, segment=0)

    labels, _ids, _ = _labels_and_ids(path, [high, low])
    assert labels == ["Zzz", "Aaa"]


def test_final_fallback_is_emission_position(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    # Identical band, segment and label: the original position decides, so the
    # emission order is preserved for genuinely indistinguishable drafts.
    first = _draft("Same", region, 0, answer="first")
    second = _draft("Same", region, 0, answer="second")

    _labels, _ids, fields = _labels_and_ids(path, [first, second])
    assert [f.answers for f in fields] == [["first"], ["second"]]


def test_primary_anchor_band_still_dominates_the_label(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    # Band 0 (label "Zzz") must precede band 1 (label "Aaa"): the label is only
    # a tie-break, it never reorders across distinct anchor bands.
    late = _draft("Aaa", region, 1)
    early = _draft("Zzz", region, 0)

    labels, _ids, _ = _labels_and_ids(path, [late, early])
    assert labels == ["Zzz", "Aaa"]


def test_untied_order_is_unchanged(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    drafts = [
        _draft("Third", region, 2),
        _draft("First", region, 0),
        _draft("Second", region, 1),
    ]
    labels, _ids, _ = _labels_and_ids(path, drafts)
    assert labels == ["First", "Second", "Third"]


# --------------------------------------------------------------------------- #
# before / after
# --------------------------------------------------------------------------- #
def test_before_after_field_id_is_stable(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    beta = _draft("Beta", region, 0)
    alpha = _draft("Alpha", region, 0)

    # BEFORE: the old stable sort keeps the emission order, so the two orders
    # produce two different orderings.
    assert _old_key_order([alpha, beta], region) == ["Alpha", "Beta"]
    assert _old_key_order([beta, alpha], region) == ["Beta", "Alpha"]

    # AFTER: the secondary key pins the order, so the field_ids agree.
    labels_a, ids_a, _ = _labels_and_ids(path, [beta, alpha])
    labels_b, ids_b, _ = _labels_and_ids(path, [alpha, beta])
    assert (labels_a, ids_a) == (labels_b, ids_b)
    # and Alpha is genuinely the tie-break winner (not a coincidence of order)
    assert labels_a == ["Alpha", "Beta"]


def test_pipeline_field_ids_are_emission_order_independent(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    def body(order):
        return json.dumps(
            {
                "fields": [
                    {
                        "label": label,
                        "control_type": "text",
                        "options": [],
                        "answer": ["v"],
                        "annotations": [],
                        "region_id": region,
                        "source_elements": [
                            {"region_id": region, "band_id": 0, "segment_index": 0}
                        ],
                    }
                    for label in order
                ]
            }
        )

    class Client:
        def __init__(self, text):
            self.text = text

        def complete(self, prompt, *, model, params):
            return LLMResponse(
                text=self.text, model=model, params=params, tokens=1, latency_ms=0
            )

    first = Pipeline(
        Store(tmp_path / "s1"), Client(body(["Beta", "Alpha"])), PipelineConfig()
    ).run(path)
    second = Pipeline(
        Store(tmp_path / "s2"), Client(body(["Alpha", "Beta"])), PipelineConfig()
    ).run(path)

    assert [(f.label_text, f.field_id) for f in first.fields] == [
        (f.label_text, f.field_id) for f in second.fields
    ]
    assert [f.label_text for f in first.fields] == ["Alpha", "Beta"]


# --------------------------------------------------------------------------- #
# the parser half: order out of parse_drafts_with_errors is irrelevant
# --------------------------------------------------------------------------- #
def test_parse_order_does_not_change_the_field_order(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    layout, _ = _layout_of(path)
    region = _label_region_id(layout)

    def field(label):
        return {
            "label": label,
            "control_type": "text",
            "options": [],
            "answer": ["v"],
            "annotations": [],
            "region_id": region,
            "source_elements": [
                {"region_id": region, "band_id": 0, "segment_index": 0}
            ],
        }

    forward, _ = parse_drafts_with_errors(json.dumps({"fields": [field("Alpha"), field("Beta")]}))
    backward, _ = parse_drafts_with_errors(json.dumps({"fields": [field("Beta"), field("Alpha")]}))

    assert [d.label for d in forward] == ["Alpha", "Beta"]
    assert [d.label for d in backward] == ["Beta", "Alpha"]

    labels_f, ids_f, _ = _labels_and_ids(path, forward)
    labels_b, ids_b, _ = _labels_and_ids(path, backward)
    assert labels_f == labels_b == ["Alpha", "Beta"]
    assert ids_f == ids_b


# --------------------------------------------------------------------------- #
# the sort key itself
# --------------------------------------------------------------------------- #
def test_sort_key_shape_beats_a_single_int():
    region = REGION
    key = _anchor_sort_key(_draft("Beta", region, 2, segment=1), region, position=7)
    assert key == (2, (2, 1), "beta", 7)

    # no refs -> the sentinel band/segment, still a tuple (not an int)
    no_refs = BindingDraft(
        draft_id="d-none",
        label="Lone",
        control_type=ControlType.TEXT,
        bbox=BBox(),
        region_id=region,
        source_refs=[],
    )
    assert _anchor_sort_key(no_refs, region, position=0) == (
        10**9,
        (10**9, 10**9),
        "lone",
        0,
    )
